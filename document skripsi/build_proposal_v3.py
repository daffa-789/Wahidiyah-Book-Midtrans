# -*- coding: utf-8 -*-
"""
PROPOSAL SKRIPSI — RANCANG BANGUN (R&D)
Universitas Wahidiyah Kediri
Program Studi Teknik Informatika — Fakultas Teknik

Mengikuti sistematika proposal "rancang bangun" Pedoman Penulisan Skripsi
Uniwa Tahun Akademik 2023-2024 (BAB II.B):
  BAB I  Pendahuluan : latar belakang, rumusan masalah, tujuan,
                       spesifikasi produk yang diharapkan, manfaat,
                       definisi operasional variabel, asumsi & keterbatasan
  BAB II Kajian Pustaka : kajian teoretis, hasil penelitian relevan,
                       kerangka konseptual, kerangka model penelitian, hipotesis
  BAB III Metode Penelitian : jenis (pengembangan), prosedur, desain uji coba,
                       subjek uji coba, jenis data, instrumen, teknik analisis data
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import json
import os

# ======================== KONSTANTA IDENTITAS ========================
NAMA     = "GHAZALAH RAMADHAN DAFFA FIRJATULLAH"
NIM      = "20232210156"
PRODI    = "PROGRAM STUDI TEKNIK INFORMATIKA"
FAKULTAS = "FAKULTAS TEKNIK"
UNIV     = "UNIVERSITAS WAHIDIYAH KEDIRI"
KOTA     = "Kediri"
TAHUN    = "2026"
GELAR    = "Sarjana Komputer (S.Kom.)"

JUDUL = ("RANCANG BANGUN PLATFORM WAHIDIYAH BOOK BERBASIS WEB: "
         "KATALOG DAN RUANG BACA DIGITAL TERINTEGRASI PEMBAYARAN QRIS "
         "DENGAN MEKANISME PENYELESAIAN TRANSAKSI IDEMPOTEN")

JUDUL_SINGKAT = "Rancang Bangun Platform Wahidiyah Book Berbasis Web"

MARGIN_KIRI  = Cm(4)
MARGIN_KANAN = Cm(3)
MARGIN_ATAS  = Cm(4)
MARGIN_BAWAH = Cm(3)
FONT_NAME = "Times New Roman"
INDENT = Cm(1.27)

HERE = os.path.dirname(os.path.abspath(__file__))
FIG_DIR = os.path.join(os.path.dirname(HERE), "docs", "diagrams")
LOCAL_LOGO = os.path.join(HERE, "_refwork", "logo_template.png")
LOGO_PATH = LOCAL_LOGO if os.path.exists(LOCAL_LOGO) else r"C:/Users/Daffa/Downloads/Aturan Skripsi/_extract/logo_template.png"
LOCAL_PUSTAKA = os.path.join(HERE, "_refwork", "daftar_pustaka_v3.json")
DATA_PUSTAKA = LOCAL_PUSTAKA if os.path.exists(LOCAL_PUSTAKA) else (
    r"C:/Users/Daffa/Downloads/Aturan Skripsi/_parsed/daftar_pustaka_v3.json"
)


# ======================== UTILITAS ========================

def set_default_font(doc):
    st = doc.styles['Normal']
    st.font.name = FONT_NAME
    st.font.size = Pt(12)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), FONT_NAME)


def set_page_margins(sec):
    sec.left_margin = MARGIN_KIRI
    sec.right_margin = MARGIN_KANAN
    sec.top_margin = MARGIN_ATAS
    sec.bottom_margin = MARGIN_BAWAH
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)


def set_run_font(run, name=FONT_NAME):
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), name)


def add_paragraph(doc, text, *, bold=False, italic=False, align=None,
                  size=12, space_before=0, space_after=6,
                  line_spacing=2.0, first_line_indent=None,
                  font_name=None, underline=False, left_indent=None,
                  style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if first_line_indent is not None:
        pf.first_line_indent = first_line_indent
    if left_indent is not None:
        pf.left_indent = left_indent
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        run.font.name = font_name or FONT_NAME
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        run.underline = underline
        set_run_font(run, font_name or FONT_NAME)
    return p


def body(doc, text, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True,
         space_after=6, size=12, italic=False, bold=False,
         line_spacing=2.0):
    return add_paragraph(doc, text, align=align, size=size,
                         first_line_indent=INDENT if indent else Cm(0),
                         space_after=space_after, italic=italic, bold=bold,
                         line_spacing=line_spacing)


def bullets(doc, items, *, level=0, numbering="number"):
    """items: list of str. numbering: 'number' | 'dash' | 'alpha'"""
    import string
    for i, it in enumerate(items):
        if numbering == "number":
            pref = f"{i+1}. "
        elif numbering == "alpha":
            pref = f"{string.ascii_lowercase[i]}. "
        else:
            pref = "\u2013 "
        add_paragraph(doc, pref + it, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                      size=12, line_spacing=2.0,
                      left_indent=Cm(1.27 + 0.6 * level),
                      first_line_indent=Cm(-0.635), space_after=4)


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_section_break(doc):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    set_page_margins(s)
    return s


def set_page_number_format(sec, fmt='lowerRoman', start=None):
    sp = sec._sectPr
    for old in sp.findall(qn('w:pgNumType')):
        sp.remove(old)
    pg = OxmlElement('w:pgNumType')
    pg.set(qn('w:fmt'), fmt)
    if start is not None:
        pg.set(qn('w:start'), str(start))
    sp.append(pg)


def add_page_number_field(paragraph):
    f1 = OxmlElement('w:fldChar')
    f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = 'PAGE   \\* MERGEFORMAT'
    f2 = OxmlElement('w:fldChar')
    f2.set(qn('w:fldCharType'), 'end')
    run = paragraph.add_run()
    run.font.size = Pt(12)
    set_run_font(run)
    r = run._element
    r.append(f1)
    r.append(it)
    r.append(f2)


def set_footer_page_number(sec):
    ft = sec.footer
    ft.is_linked_to_previous = False
    for p in list(ft.paragraphs):
        if p.text or len(p.runs) > 0:
            p.clear()
    p = ft.paragraphs[0] if ft.paragraphs else ft.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number_field(p)


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), r_id)
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    c = OxmlElement('w:color')
    c.set(qn('w:val'), '0563C1')
    rPr.append(c)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)
    rf = OxmlElement('w:rFonts')
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rf.set(qn(a), FONT_NAME)
    rPr.append(rf)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '24')
    rPr.append(sz)
    r.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def add_centered_image(doc, path, width_cm):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(path, width=Cm(width_cm))
    return p


# ======================== GAYA JUDUL (Heading 1/2/3) ========================

def _init_heading_styles(doc):
    spec = {
        'Heading 1': dict(size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                          before=12, after=18, spacing=1.5),
        'Heading 2': dict(size=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT,
                          before=12, after=6, spacing=1.5),
        'Heading 3': dict(size=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT,
                          before=8, after=4, spacing=1.5),
    }
    for name, s in spec.items():
        st = doc.styles[name]
        st.font.name = FONT_NAME
        st.font.size = Pt(s['size'])
        st.bold = s['bold']
        st.italic = False
        st.underline = False
        st.font.all_caps = False
        st.font.color.rgb = RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            rpr.insert(0, rf)
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(a), FONT_NAME)
        pf = st.paragraph_format
        pf.alignment = s['align']
        pf.space_before = Pt(s['before'])
        pf.space_after = Pt(s['after'])
        pf.line_spacing = s['spacing']
        pf.keep_with_next = True
        pf.first_line_indent = Cm(0)
        pf.left_indent = Cm(0)


def h1(doc, *lines):
    p = doc.add_paragraph(style='Heading 1')
    for i, ln in enumerate(lines):
        if i:
            p.add_run().add_break()
        r = p.add_run(ln)
        r.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(0, 0, 0)
        set_run_font(r)
    return p


def h2(doc, text):
    p = doc.add_paragraph(style='Heading 2')
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0, 0, 0)
    set_run_font(r)
    return p


def h3(doc, text):
    p = doc.add_paragraph(style='Heading 3')
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0, 0, 0)
    set_run_font(r)
    return p


def bab(doc, roman, title):
    return h1(doc, f"BAB {roman}", title)


# ======================== TABEL & KETERANGAN ========================

def add_table(doc, headers, rows, widths=None, fs=10, hdr_fs=10):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.autofit = True
    hdr = t.rows[0].cells
    for i, htxt in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(htxt)
        r.bold = True
        r.font.size = Pt(hdr_fs)
        set_run_font(r)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(2)
            p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if i == 0
                           else WD_ALIGN_PARAGRAPH.JUSTIFY)
            r = p.add_run(str(val))
            r.font.size = Pt(fs)
            set_run_font(r)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t


def caption(doc, text, *, center=True, before=8, after=12, fs=11):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    r.font.size = Pt(fs)
    r.bold = True
    set_run_font(r)
    return p


def figure(doc, filename, caption_text, width_cm=14.0):
    path = os.path.join(FIG_DIR, filename)
    if os.path.exists(path):
        add_centered_image(doc, path, width_cm)
    caption(doc, caption_text)
    add_paragraph(doc, "", size=6, line_spacing=1.0, space_after=6)


def code_block(doc, lines):
    for ln in lines:
        add_paragraph(doc, ln if ln else " ", size=9, font_name="Consolas",
                      line_spacing=1.0, space_after=0,
                      left_indent=Cm(1.0))


# ======================== HALAMAN AWAL ========================

def build_cover(doc):
    add_paragraph(doc, UNIV, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                  size=14, line_spacing=1.5, space_before=18, space_after=4)
    add_paragraph(doc, FAKULTAS, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                  size=13, line_spacing=1.5, space_after=2)
    add_paragraph(doc, PRODI, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                  size=13, line_spacing=1.5, space_after=16)

    if os.path.exists(LOGO_PATH):
        add_centered_image(doc, LOGO_PATH, 5)

    add_paragraph(doc, JUDUL, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                  size=14, line_spacing=1.0, space_before=18, space_after=18)
    add_paragraph(doc, "PROPOSAL PENELITIAN", bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, size=14,
                  line_spacing=1.0, space_after=18)
    add_paragraph(doc, "Oleh:", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                  size=12, line_spacing=1.0, space_after=8)
    add_paragraph(doc, NAMA, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                  size=13, line_spacing=1.0, space_after=4)
    add_paragraph(doc, f"NIM. {NIM}", align=WD_ALIGN_PARAGRAPH.CENTER,
                  size=12, line_spacing=1.0, space_after=24)
    add_paragraph(doc, KOTA.upper(), align=WD_ALIGN_PARAGRAPH.CENTER,
                  size=12, line_spacing=1.0, space_before=24, space_after=2)
    add_paragraph(doc, TAHUN, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                  size=12, line_spacing=1.0, space_after=2)


def build_halaman_judul(doc):
    add_paragraph(doc, JUDUL, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                  size=14, line_spacing=1.0, space_before=12, space_after=18)
    add_paragraph(doc, "PROPOSAL PENELITIAN", bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, size=14,
                  line_spacing=1.0, space_after=12)
    add_paragraph(doc, f"Diajukan kepada Program Sarjana {UNIV}",
                  align=WD_ALIGN_PARAGRAPH.CENTER, size=12,
                  line_spacing=1.0, space_after=2)
    add_paragraph(doc, "untuk Memenuhi Sebagian Persyaratan dalam Memperoleh "
                       f"Gelar {GELAR}", align=WD_ALIGN_PARAGRAPH.CENTER,
                  size=12, line_spacing=1.0, space_after=2)
    add_paragraph(doc, f"pada {PRODI}", align=WD_ALIGN_PARAGRAPH.CENTER,
                  size=12, line_spacing=1.0, space_after=18)
    add_paragraph(doc, "Oleh:", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                  size=12, line_spacing=1.0, space_after=8)
    add_paragraph(doc, NAMA, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                  size=13, line_spacing=1.0, space_after=4)
    add_paragraph(doc, f"NIM. {NIM}", align=WD_ALIGN_PARAGRAPH.CENTER,
                  size=12, line_spacing=1.0, space_after=24)
    for t in (UNIV, FAKULTAS, PRODI, TAHUN):
        add_paragraph(doc, t, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                      size=12, line_spacing=1.0, space_after=2)


def build_persetujuan(doc):
    h1(doc, "HALAMAN PERSETUJUAN")
    body(doc, f"Proposal penelitian oleh {NAMA}, NIM. {NIM}, dengan judul:",
         align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, space_after=6)
    add_paragraph(doc, f"\u201C{JUDUL}\u201D", bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, size=12,
                  line_spacing=2.0, space_after=12)
    body(doc, "telah diperiksa, disetujui, dan layak untuk diujikan dalam "
              "seminar proposal penelitian.", align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         indent=True, space_after=18)

    add_table(doc,
              ["Jabatan", "Nama", "Tanda Tangan", "Tanggal"],
              [["Pembimbing I", "", "", ""],
               ["Pembimbing II", "", "", ""],
               ["Ketua Program Studi", "", "", ""]],
              widths=[4.5, 4.5, 3.0, 2.0], fs=11)
    add_paragraph(doc, "", size=12, line_spacing=1.0, space_after=12)
    add_paragraph(doc, f"{KOTA}, ................................ {TAHUN}",
                  align=WD_ALIGN_PARAGRAPH.RIGHT, line_spacing=1.5,
                  space_after=6)
    add_paragraph(doc, "Mengetahui,", align=WD_ALIGN_PARAGRAPH.RIGHT,
                  line_spacing=1.5, space_after=48)
    add_paragraph(doc, "Ketua Program Studi Teknik Informatika", bold=True,
                  align=WD_ALIGN_PARAGRAPH.RIGHT, line_spacing=1.5,
                  space_after=2)


def build_pengesahan(doc):
    h1(doc, "HALAMAN PENGESAHAN")
    body(doc, f"Proposal penelitian dengan judul \u201C{JUDUL_SINGKAT}\u201D "
              "yang disusun oleh:", align=WD_ALIGN_PARAGRAPH.LEFT,
         indent=False, space_after=8)
    add_paragraph(doc, f"Nama\t\t: {NAMA}", line_spacing=2.0,
                  first_line_indent=Cm(1.0), space_after=2)
    add_paragraph(doc, f"NIM\t\t: {NIM}", line_spacing=2.0,
                  first_line_indent=Cm(1.0), space_after=2)
    add_paragraph(doc, f"Program Studi\t: Teknik Informatika",
                  line_spacing=2.0, first_line_indent=Cm(1.0), space_after=2)
    add_paragraph(doc, f"Fakultas\t\t: {FAKULTAS.title()}",
                  line_spacing=2.0, first_line_indent=Cm(1.0), space_after=12)
    body(doc, "telah diujikan dalam seminar proposal dan dinyatakan "
              "DITERIMA / DITOLAK*) pada:", align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         indent=True, space_after=12)

    add_table(doc,
              ["No.", "Nama Penguji", "Jabatan", "Tanda Tangan"],
              [["1.", "", "Penguji I", ""],
               ["2.", "", "Penguji II", ""],
               ["3.", "", "Penguji III", ""]],
              widths=[1.2, 5.0, 4.3, 3.5], fs=11)
    add_paragraph(doc, "", size=12, line_spacing=1.0, space_after=8)
    add_paragraph(doc, "*) Coret yang tidak perlu", italic=True, size=11,
                  line_spacing=1.5, space_after=12)
    add_paragraph(doc, f"{KOTA}, ................................ {TAHUN}",
                  align=WD_ALIGN_PARAGRAPH.RIGHT, line_spacing=1.5,
                  space_after=6)
    add_paragraph(doc, "Dekan Fakultas Teknik", bold=True,
                  align=WD_ALIGN_PARAGRAPH.RIGHT, line_spacing=1.5,
                  space_after=48)
    add_paragraph(doc, "(..................................................)",
                  align=WD_ALIGN_PARAGRAPH.RIGHT, line_spacing=1.0,
                  space_after=2)


def build_surat_pernyataan(doc):
    h1(doc, "SURAT PERNYATAAN KEASLIAN")
    body(doc, "Yang bertanda tangan di bawah ini:",
         align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, space_after=6)
    add_paragraph(doc, f"Nama\t\t: {NAMA}", line_spacing=2.0,
                  first_line_indent=Cm(1.0), space_after=2)
    add_paragraph(doc, f"NIM\t\t: {NIM}", line_spacing=2.0,
                  first_line_indent=Cm(1.0), space_after=2)
    add_paragraph(doc, "Program Studi\t: Teknik Informatika",
                  line_spacing=2.0, first_line_indent=Cm(1.0), space_after=2)
    add_paragraph(doc, f"Fakultas\t\t: {FAKULTAS.title()}",
                  line_spacing=2.0, first_line_indent=Cm(1.0), space_after=12)
    body(doc, "Dengan ini menyatakan dengan sesungguhnya bahwa proposal "
              "penelitian yang berjudul:")
    add_paragraph(doc, f"\u201C{JUDUL}\u201D", bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, size=12,
                  line_spacing=2.0, space_after=12)
    body(doc, "adalah benar-benar karya saya sendiri, bukan hasil "
              "penjiplakan (plagiasi) dari karya orang lain, dan belum pernah "
              "diajukan untuk memperoleh gelar akademik di perguruan tinggi "
              "mana pun. Apabila di kemudian hari terbukti bahwa pernyataan "
              "ini tidak benar, saya bersedia menerima sanksi akademik "
              "sesuai dengan peraturan yang berlaku di Universitas Wahidiyah "
              "Kediri.")
    add_paragraph(doc, "", size=12, line_spacing=1.0, space_after=12)
    add_paragraph(doc, f"{KOTA}, ................................ {TAHUN}",
                  align=WD_ALIGN_PARAGRAPH.RIGHT, line_spacing=1.5,
                  space_after=6)
    add_paragraph(doc, "Yang Menyatakan,", align=WD_ALIGN_PARAGRAPH.RIGHT,
                  line_spacing=1.5, space_after=48)
    add_paragraph(doc, NAMA, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT,
                  line_spacing=1.0, space_after=2)
    add_paragraph(doc, f"NIM. {NIM}", align=WD_ALIGN_PARAGRAPH.RIGHT,
                  line_spacing=1.0, space_after=6)


def build_kata_pengantar(doc):
    h1(doc, "KATA PENGANTAR")
    body(doc, "Puji syukur kehadirat Allah SWT atas rahmat, taufik, dan "
              "hidayah-Nya sehingga penulis dapat menyelesaikan penyusunan "
              "proposal penelitian yang berjudul \u201C" + JUDUL_SINGKAT +
              "\u201D. Proposal ini disusun sebagai salah satu persyaratan "
              "untuk menempuh seminar proposal penelitian pada Program Studi "
              "Teknik Informatika, Fakultas Teknik, Universitas Wahidiyah "
              "Kediri.")
    body(doc, "Penulis menyadari bahwa penyusunan proposal ini tidak lepas "
              "dari bantuan, bimbingan, dan dukungan berbagai pihak. Oleh "
              "karena itu, dengan segala kerendahan hati penulis "
              "menyampaikan terima kasih yang sebesar-besarnya kepada:")
    bullets(doc, [
        "Rektor Universitas Wahidiyah Kediri beserta jajaran pimpinan "
        "universitas yang telah memberikan kesempatan dan fasilitas kepada "
        "penulis untuk menempuh pendidikan.",
        "Dekan Fakultas Teknik Universitas Wahidiyah Kediri beserta staf "
        "yang telah memberikan dukungan akademik dan administratif.",
        "Ketua Program Studi Teknik Informatika Universitas Wahidiyah Kediri "
        "yang telah memberikan arahan dan bimbingan akademik.",
        "Dosen Pembimbing yang dengan sabar telah meluangkan waktu, "
        "memberikan bimbingan, saran, dan koreksi yang sangat berharga dalam "
        "penyusunan proposal ini.",
        "Seluruh dosen Program Studi Teknik Informatika yang telah "
        "membekali penulis dengan ilmu pengetahuan dan keterampilan di "
        "bidang rekayasa perangkat lunak.",
        "Pengelola dan seluruh pengguna platform Wahidiyah Book yang telah "
        "bersedia menjadi narasumber dan subjek uji coba dalam penelitian "
        "ini.",
        "Kedua orang tua dan keluarga tercinta yang senantiasa memberikan "
        "doa, dukungan moral, dan semangat yang tidak pernah putus.",
        "Rekan-rekan mahasiswa Teknik Informatika Universitas Wahidiyah "
        "Kediri yang telah memberikan bantuan, masukan, dan kebersamaan "
        "selama proses penyusunan proposal ini.",
    ], numbering="number")
    body(doc, "Penulis menyadari sepenuhnya bahwa proposal ini masih jauh "
              "dari kata sempurna. Oleh karena itu, kritik dan saran yang "
              "membangun dari segala pihak sangat penulis harapkan demi "
              "kesempurnaan proposal ini. Semoga proposal ini dapat "
              "memberikan manfaat bagi pengembangan ilmu pengetahuan, "
              "khususnya di bidang rekayasa perangkat lunak, serta bagi "
              "pengembangan literasi digital di lingkungan Universitas "
              "Wahidiyah Kediri.")
    add_paragraph(doc, "", size=12, line_spacing=1.0, space_after=12)
    add_paragraph(doc, f"{KOTA}, ................................ {TAHUN}",
                  align=WD_ALIGN_PARAGRAPH.RIGHT, line_spacing=1.5,
                  space_after=6)
    add_paragraph(doc, "Penulis,", align=WD_ALIGN_PARAGRAPH.RIGHT,
                  line_spacing=1.5, space_after=48)
    add_paragraph(doc, NAMA, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT,
                  line_spacing=1.0, space_after=2)
    add_paragraph(doc, f"NIM. {NIM}", align=WD_ALIGN_PARAGRAPH.RIGHT,
                  line_spacing=1.0, space_after=6)


def build_abstrak(doc):
    h1(doc, "ABSTRAK")
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(NAMA + ", " + NIM + ". ")
    r.font.size = Pt(12)
    set_run_font(r)
    r = p.add_run(JUDUL_SINGKAT + ": Katalog dan Ruang Baca Digital "
                  "Terintegrasi Pembayaran QRIS dengan Mekanisme Penyelesaian "
                  "Transaksi Idempoten. ")
    r.bold = True
    r.font.size = Pt(12)
    set_run_font(r)
    r = p.add_run(f"Proposal Penelitian, Program Studi Teknik Informatika, "
                  f"{FAKULTAS.title()}, {UNIV.title()}, {TAHUN}.")
    r.italic = True
    r.font.size = Pt(12)
    set_run_font(r)

    body(doc, "Pengelolaan katalog buku dan layanan ruang baca di lingkungan "
              "Universitas Wahidiyah Kediri masih dilakukan secara "
              "terpisah-pisah dan belum terintegrasi dengan sistem "
              "pembayaran digital, sehingga pengguna kesulitan menemukan "
              "bahan bacaan dan pengelola tidak memiliki data transaksi yang "
              "tercatat rapi. Penelitian ini bertujuan merancang dan membangun "
              "platform Wahidiyah Book berbasis web yang memadukan katalog "
              "buku, ruang baca digital berlangganan, dan pembayaran QRIS "
              "dalam satu kesatuan sistem. Permasalahan teknis yang diangkat "
              "bukan sekadar integrasi payment gateway, melainkan konsistensi "
              "penyelesaian transaksi: notifikasi Midtrans dapat dikirim "
              "berulang kali dan dapat hilang sama sekali, sementara "
              "pengembangan lokal tidak memiliki URL publik untuk menerima "
              "notifikasi. Untuk itu dirancang mekanisme penyelesaian "
              "transaksi yang idempoten dengan pola compare-and-swap "
              "atomik \u2014 pembaruan bersyarat UPDATE \u2026 WHERE id = ? "
              "AND status = 'pending' \u2014 yang dikombinasikan dengan "
              "transaksi kompensasi apabila aktivasi langganan gagal, serta "
              "dua jalur penyelesaian (webhook dan polling yang otoritatif di "
              "sisi server) yang bermuara pada fungsi penyelesaian yang sama. "
              "Keamanan notifikasi diterapkan secara fail-closed melalui "
              "verifikasi tanda tangan SHA-512 dengan perbandingan waktu "
              "konstan (timingSafeEqual) dan rekonsiliasi nominal, "
              "sedangkan autentikasi menggunakan pendaftaran dua fase "
              "berbasis OTP yang baru membuat baris pengguna setelah bukti OTP "
              "sah. Metode yang digunakan adalah Research and Development "
              "(rancang bangun) dengan model prototyping yang meliputi "
              "analisis kebutuhan, perancangan, implementasi, pengujian "
              "internal, uji coba pengguna, dan evaluasi iteratif. Produk "
              "diuji dengan pengujian black box berbasis ISO/IEC 25010, "
              "eksperimen idempotensi dengan pengiriman notifikasi ganda, "
              "serta kuesioner System Usability Scale (SUS). Luaran yang "
              "diharapkan berupa platform yang berfungsi penuh, bebas dari "
              "aktivasi langganan ganda, dan memenuhi ambang keberterimaan "
              "usability.")
    add_paragraph(doc, "", size=12, line_spacing=1.0, space_after=6)
    add_paragraph(doc, "Kata Kunci: rancang bangun, Wahidiyah Book, QRIS, "
                       "idempotensi, compare-and-swap, progressive web app, "
                       "ISO/IEC 25010, System Usability Scale",
                  align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12,
                  line_spacing=1.5, first_line_indent=INDENT, italic=True)


# ======================== DAFTAR ISI / TABEL / GAMBAR ========================

def add_toc_field(doc):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run()
    f1 = OxmlElement('w:fldChar')
    f1.set(qn('w:fldCharType'), 'begin')
    f1.set(qn('w:dirty'), 'true')
    it = OxmlElement('w:instrText')
    it.set(qn('xml:space'), 'preserve')
    it.text = 'TOC \\o "1-3" \\h \\z \\u'
    f2 = OxmlElement('w:fldChar')
    f2.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t')
    t.text = ('Buka dokumen ini di Microsoft Word, lalu tekan Ctrl+A dan F9 '
              '(atau klik kanan \u2192 Update Field) untuk mengisi nomor halaman '
              'daftar isi secara otomatis.')
    f3 = OxmlElement('w:fldChar')
    f3.set(qn('w:fldCharType'), 'end')
    el = r._element
    el.append(f1)
    el.append(it)
    el.append(f2)
    el.append(t)
    el.append(f3)


def build_daftar_isi(doc):
    h1(doc, "DAFTAR ISI")
    add_toc_field(doc)


def build_daftar_tabel(doc):
    h1(doc, "DAFTAR TABEL")
    rows = [
        ["Tabel 1.1", "Spesifikasi non-fungsional produk berdasarkan "
                      "karakteristik ISO/IEC 25010"],
        ["Tabel 1.2", "Definisi operasional variabel penelitian"],
        ["Tabel 2.1", "Hasil penelitian terdahulu yang relevan"],
        ["Tabel 2.2", "Kerangka model penelitian rancang bangun Wahidiyah Book"],
        ["Tabel 3.1", "Susunan teknologi implementasi"],
        ["Tabel 3.2", "Contoh skenario pengujian black box"],
        ["Tabel 3.3", "Subjek uji coba penelitian"],
        ["Tabel 3.4", "Instrumen pengumpulan data"],
        ["Tabel 3.5", "Kriteria penilaian skor System Usability Scale (SUS)"],
        ["Tabel 3.6", "Jadwal pelaksanaan penelitian"],
    ]
    add_table(doc, ["No.", "Uraian"], rows, widths=[2.5, 11.5], fs=11)


def build_daftar_gambar(doc):
    h1(doc, "DAFTAR GAMBAR")
    rows = [
        ["Gambar 2.1", "Arsitektur sistem tiga lapis platform Wahidiyah Book"],
        ["Gambar 2.2", "Alur penyelesaian transaksi QRIS yang idempoten "
                       "dengan compare-and-swap dan kompensasi"],
        ["Gambar 2.3", "Dua jalur penyelesaian transaksi: webhook dan polling "
                       "otoritatif server"],
        ["Gambar 2.4", "Registrasi dua fase berbasis OTP"],
        ["Gambar 2.5", "Kerangka konseptual penelitian"],
        ["Gambar 3.1", "Tahapan prosedur pengembangan model prototyping"],
    ]
    add_table(doc, ["No.", "Uraian"], rows, widths=[2.5, 11.5], fs=11)


# ======================== BAB I ========================

def build_bab1(doc):
    bab(doc, "I", "PENDAHULUAN")

    # ---------- A ----------
    h2(doc, "A. Latar Belakang Masalah")
    body(doc, "Perkembangan teknologi informasi dan komunikasi telah "
              "mengubah cara manusia mengakses, mengelola, dan "
              "mendistribusikan pengetahuan. Di lingkungan perguruan tinggi, "
              "perpustakaan tidak lagi berfungsi semata sebagai ruang fisik "
              "penyimpanan buku, melainkan bertransformasi menjadi perpustakaan "
              "digital yang menyediakan katalog daring, ruang baca digital, "
              "dan layanan sirkulasi yang dapat diakses tanpa batas waktu. "
              "Transformasi ini menuntut tersedianya perangkat lunak yang "
              "mampu mengintegrasikan katalog, hak akses konten, dan transaksi "
              "dalam satu platform.")
    body(doc, "Universitas Wahidiyah Kediri merupakan salah satu perguruan "
              "tinggi yang terus berupaya meningkatkan mutu layanan akademik "
              "bagi civitas akademikanya. Namun demikian, pengelolaan bahan "
              "bacaan di lingkungan universitas masih menghadapi beberapa "
              "kendala. Pertama, informasi mengenai koleksi buku belum "
              "terhimpun dalam satu katalog digital yang mudah ditelusuri, "
              "sehingga pengguna kesulitan menemukan judul yang dibutuhkan. "
              "Kedua, belum tersedia ruang baca digital yang mampu membedakan "
              "konten terbuka dan konten berbayar secara konsisten. Ketiga, "
              "apabila layanan berlangganan disediakan, belum ada mekanisme "
              "pembayaran digital yang tercatat otomatis, padahal QRIS (Quick "
              "Response Code Indonesia Standard) telah menjadi standar "
              "pembayaran nasional yang digunakan secara luas di Indonesia "
              "(Rachman dkk., 2024; Gea & Al-Azhar, 2021).")
    body(doc, "Berbagai penelitian rancang bangun sistem informasi di "
              "bidang perpustakaan digital dan e-commerce kampus telah "
              "dilakukan, antara lain oleh Mulyani dan Hati (2021), Saekup "
              "(2022), serta Wijana dan Prasetio (2023). Penelitian-penelitian "
              "tersebut umumnya berfokus pada digitalisasi katalog dan "
              "sirkulasi. Di sisi lain, penelitian mengenai integrasi payment "
              "gateway seperti yang dilakukan oleh Muttaqin dkk. (2022), "
              "Nariswari dkk. (2023), Rezkiandra (2024), dan Richo dkk. (2021) "
              "pada umumnya berhenti pada pernyataan bahwa integrasi "
              "pembayaran telah berhasil dilakukan, tanpa menguji bagaimana "
              "sistem menjaga konsistensi status transaksi apabila notifikasi "
              "dari penyedia pembayaran datang lebih dari satu kali.")
    body(doc, "Padahal, pada kenyataannya notifikasi pembayaran bersifat "
              "tidak dapat diandalkan (unreliable delivery). Penyedia "
              "pembayaran lazim melakukan pengiriman ulang (retry) apabila "
              "titik akhir penerima tidak merespons dengan cepat, dan "
              "notifikasi dapat hilang sama sekali apabila jaringan "
              "bermasalah atau aplikasi sedang tidak berjalan. Apabila "
              "penanganan notifikasi tidak dirancang dengan benar, dua "
              "kemungkinan kegagalan dapat terjadi: pertama, langganan "
              "pengguna diaktivasi lebih dari satu kali sehingga terjadi "
              "penggandaan masa berlaku dan kerugian bagi pengelola; kedua, "
              "pembayaran yang sebenarnya sudah lunas tidak pernah tercatat "
              "sehingga pengguna tidak memperoleh hak aksesnya. Kedua "
              "kegagalan ini bermuara pada satu akar masalah yang sama, yaitu "
              "tidak idempotennya operasi penyelesaian transaksi.")
    body(doc, "Selain persoalan tersebut, terdapat kendala praktis dalam "
              "pengembangan sistem pembayaran di lingkungan kampus: server "
              "pengembangan yang dijalankan pada localhost tidak memiliki URL "
              "publik, sehingga penyedia pembayaran tidak dapat mengirimkan "
              "notifikasi ke alamat tersebut. Kondisi ini membuat pengujian "
              "alur pembayaran menjadi sulit dan sering kali mendorong "
              "pengembang membuat jalan pintas yang tidak aman, misalnya "
              "menganggap pembayaran berhasil hanya berdasarkan permintaan "
              "dari sisi klien.")
    body(doc, "Berdasarkan uraian di atas, penelitian ini mengusulkan "
              "rancang bangun platform Wahidiyah Book berbasis web yang "
              "mengintegrasikan katalog buku, ruang baca digital berlangganan, "
              "dan pembayaran QRIS. Kontribusi yang ingin dicapai bukan hanya "
              "tersedianya produk, melainkan rancangan mekanisme penyelesaian "
              "transaksi yang idempoten, terverifikasi secara fail-closed, dan "
              "tetap dapat diuji dalam kondisi pengembangan lokal. "
              "Perancangan ini memanfaatkan pola compare-and-swap atomik pada "
              "tingkat basis data, transaksi kompensasi, verifikasi tanda "
              "tangan SHA-512 dengan perbandingan waktu konstan, serta dua "
              "jalur penyelesaian transaksi yang bermuara pada satu fungsi "
              "yang sama. Dengan demikian, penelitian ini diharapkan "
              "menghasilkan produk yang tidak sekadar berfungsi, tetapi juga "
              "memiliki justifikasi rekayasa yang dapat dipertanggungjawabkan.")

    # ---------- B ----------
    h2(doc, "B. Rumusan Masalah")
    body(doc, "Berdasarkan latar belakang tersebut, rumusan masalah dalam "
              "penelitian ini adalah sebagai berikut.", indent=True)
    bullets(doc, [
        "Bagaimana merancang dan membangun platform Wahidiyah Book berbasis "
        "web yang mengintegrasikan katalog buku, ruang baca digital, dan "
        "layanan langganan berbayar?",
        "Bagaimana merancang mekanisme penyelesaian transaksi QRIS yang "
        "idempoten sehingga notifikasi ganda maupun penyelesaian yang "
        "berbarengan tidak menyebabkan aktivasi langganan ganda?",
        "Bagaimana merancang verifikasi notifikasi pembayaran yang bersifat "
        "fail-closed melalui tanda tangan SHA-512, perbandingan waktu "
        "konstan, dan rekonsiliasi nominal?",
        "Bagaimana merancang dua jalur penyelesaian transaksi, yaitu "
        "notifikasi (webhook) dan pemeriksaan status yang otoritatif di sisi "
        "server, sehingga sistem tetap dapat diuji dan beroperasi tanpa URL "
        "publik?",
        "Bagaimana merancang autentikasi pendaftaran dua fase berbasis OTP "
        "sehingga tidak terbentuk akun yang belum terverifikasi di basis data?",
        "Bagaimana merancang strategi tembolok (cache) pada progressive web "
        "app agar konten berbayar, data autentikasi, dan data transaksi tidak "
        "tersimpan di perangkat?",
        "Bagaimana kualitas fungsional dan tingkat kebergunaan (usability) "
        "platform yang dihasilkan apabila diukur dengan pengujian black box "
        "berbasis ISO/IEC 25010 dan kuesioner System Usability Scale (SUS)?",
    ], numbering="number")

    # ---------- C ----------
    h2(doc, "C. Tujuan Penelitian")
    body(doc, "Tujuan yang ingin dicapai dalam penelitian ini adalah sebagai "
              "berikut.", indent=True)
    bullets(doc, [
        "Merancang dan membangun platform Wahidiyah Book berbasis web yang "
        "mengintegrasikan katalog buku, ruang baca digital, dan layanan "
        "langganan berbayar.",
        "Merancang dan mengimplementasikan mekanisme penyelesaian transaksi "
        "QRIS yang idempoten berbasis compare-and-swap atomik dan transaksi "
        "kompensasi sehingga aktivasi langganan ganda dapat dicegah.",
        "Mengimplementasikan verifikasi notifikasi pembayaran yang fail-closed "
        "menggunakan tanda tangan SHA-512, perbandingan waktu konstan "
        "(timingSafeEqual), dan rekonsiliasi nominal transaksi.",
        "Menyediakan dua jalur penyelesaian transaksi, yaitu webhook dan "
        "pemeriksaan status otoritatif di sisi server, agar proses "
        "pengembangan dan pengujian dapat dilakukan tanpa URL publik.",
        "Merancang autentikasi pendaftaran dua fase berbasis OTP yang hanya "
        "membuat baris pengguna setelah kode OTP terbukti sah.",
        "Merancang strategi tembolok selektif pada service worker sehingga "
        "konten berbayar dan data sensitif tidak pernah disimpan di perangkat.",
        "Mengukur kualitas fungsional dan tingkat kebergunaan platform yang "
        "dihasilkan dengan pengujian black box berbasis ISO/IEC 25010 dan "
        "kuesioner SUS.",
    ], numbering="number")

    # ---------- D ----------
    h2(doc, "D. Spesifikasi Produk yang Diharapkan")
    body(doc, "Produk yang dihasilkan dalam penelitian ini adalah sebuah "
              "platform web bernama Wahidiyah Book. Spesifikasi fungsional "
              "produk yang diharapkan adalah sebagai berikut.", indent=True)
    bullets(doc, [
        "Dapat diakses melalui peramban pada perangkat komputer maupun "
        "telepon pintar, serta dapat dipasang sebagai progressive web app "
        "(PWA) dengan cangkang aplikasi (app shell) yang tetap tersedia saat "
        "jaringan terputus.",
        "Menyediakan modul katalog buku yang menampilkan judul, penulis, "
        "sinopsis, dan sampul yang telah dioptimasi, lengkap dengan fungsi "
        "pencarian dan penyaringan.",
        "Menyediakan ruang baca digital dengan pembatasan akses berbasis "
        "status langganan: buku terbuka dapat dibaca oleh siapa pun, "
        "sedangkan buku berbayar hanya dapat dibaca oleh pengguna dengan "
        "langganan aktif.",
        "Menyediakan alur transaksi QRIS yang mencakup pembuatan kode QR, "
        "batas waktu kedaluwarsa, dan pembaruan status transaksi secara "
        "otomatis.",
        "Menyelesaikan transaksi secara idempoten: pengiriman ulang "
        "notifikasi maupun penyelesaian berbarengan dari dua jalur yang "
        "berbeda tidak menghasilkan aktivasi langganan ganda.",
        "Menyediakan autentikasi pendaftaran dua fase berbasis OTP dengan "
        "batas waktu 10 menit dan pembatasan percobaan, serta masuk "
        "(login) menggunakan surel dan kata sandi maupun akun Google.",
        "Menyediakan dasbor pengelola (admin) yang memuat ringkasan "
        "statistik, pengelolaan buku, pengelolaan peran pengguna, agenda, "
        "karousel beranda, dan log aktivitas.",
        "Menyediakan pengelolaan langganan otomatis, termasuk penonaktifan "
        "langganan yang telah melewati masa berlaku.",
        "Menerapkan pengamanan dasar: tajuk keamanan (security headers), "
        "pembatasan laju permintaan (rate limiting), validasi masukan, "
        "penyimpanan kata sandi dengan bcrypt, dan pembatasan akses data pada "
        "tingkat baris (row level security).",
    ], numbering="number")

    body(doc, "Adapun spesifikasi non-fungsional produk ditetapkan "
              "berdasarkan karakteristik kualitas perangkat lunak ISO/IEC "
              "25010 yang disajikan pada Tabel 1.1.", indent=True,
         space_after=8)

    add_table(doc,
              ["Karakteristik ISO/IEC 25010", "Target yang Diharapkan"],
              [["Kesesuaian fungsional\n(functional suitability)",
                "Seluruh skenario pengujian black box pada fungsi utama "
                "berstatus valid (100%)."],
               ["Keandalan\n(reliability)",
                "Penyelesaian transaksi idempoten: tidak ada aktivasi "
                "langganan ganda pada 10 pengiriman ulang notifikasi."],
               ["Kegunaan\n(usability)",
                "Skor SUS \u2265 68 (kategori acceptable) dengan "
                "tingkat penyelesaian tugas \u2265 90%."],
               ["Efisiensi kinerja\n(performance efficiency)",
                "Waktu muat halaman katalog < 3 detik pada koneksi "
                "seluler; berkas gambar dikonversi ke format WebP."],
               ["Keamanan\n(security)",
                "Notifikasi tanpa tanda tangan sah ditolak; konten berbayar "
                "dan data autentikasi tidak tersimpan di tembolok PWA."],
               ["Kemudahan pemeliharaan\n(maintainability)",
                "Kode sumber terstruktur berlapis (rute, librari, "
                "middleware) dan dapat dibangun ulang dengan satu perintah."],
               ["Portabilitas\n(portability)",
                "Dapat dijalankan pada sistem operasi Windows, Linux, dan "
                "macOS melalui Node.js."]],
              widths=[5.5, 8.5], fs=10)
    caption(doc, "Tabel 1.1 Spesifikasi non-fungsional produk berdasarkan "
                 "karakteristik ISO/IEC 25010")

    # ---------- E ----------
    h2(doc, "E. Manfaat Penelitian")
    h3(doc, "1. Manfaat Teoretis")
    body(doc, "Penelitian ini diharapkan memberikan kontribusi bagi "
              "pengembangan keilmuan di bidang rekayasa perangkat lunak, "
              "khususnya mengenai penerapan pola compare-and-swap dan "
              "transaksi kompensasi untuk menjaga idempotensi penyelesaian "
              "transaksi pada sistem yang terintegrasi dengan payment "
              "gateway. Hasil penelitian ini juga dapat menjadi bahan "
              "rujukan bagi penelitian rancang bangun sejenis yang "
              "membutuhkan justifikasi mengenai konsistensi status transaksi "
              "dan keamanan notifikasi.")
    h3(doc, "2. Manfaat Praktis")
    bullets(doc, [
        "Bagi Universitas Wahidiyah Kediri, menyediakan platform terpadu "
        "untuk mengelola katalog bahan bacaan dan layanan ruang baca digital, "
        "serta memperoleh data transaksi yang tercatat rapi.",
        "Bagi pengelola (admin), mempermudah pemantauan aktivitas pengguna, "
        "pengelolaan konten, dan penertiban hak akses konten berbayar.",
        "Bagi pengguna, memberikan kemudahan dalam menemukan, membaca, dan "
        "berlangganan bahan bacaan dengan pembayaran QRIS yang aman dan "
        "langsung tercatat.",
        "Bagi peneliti, memberikan pengalaman rekayasa langsung dalam "
        "merancang sistem yang tangguh terhadap kegagalan jaringan dan "
        "pengiriman pesan ganda.",
    ], numbering="number")

    # ---------- F ----------
    h2(doc, "F. Definisi Operasional Variabel")
    body(doc, "Agar tidak terjadi perbedaan penafsiran, variabel-variabel "
              "yang digunakan dalam penelitian ini didefinisikan secara "
              "operasional sebagaimana disajikan pada Tabel 1.2.",
         indent=True, space_after=8)
    add_table(doc,
              ["Variabel", "Definisi Operasional", "Indikator / Cara Ukur"],
              [["Penyelesaian transaksi idempoten",
                "Sifat operasi penyelesaian transaksi QRIS yang menghasilkan "
                "keadaan akhir yang sama berapa pun kali operasi tersebut "
                "dijalankan untuk satu order_id yang sama.",
                "Jumlah aktivasi langganan yang tercatat pada tabel "
                "subscriptions untuk satu order_id. Idempoten apabila "
                "berjumlah tepat 1."],
               ["Rasio aktivasi ganda",
                "Perbandingan antara jumlah order_id yang menghasilkan lebih "
                "dari satu aktivasi dengan seluruh order_id yang diuji.",
                "(jumlah aktivasi ganda / jumlah order_id) \u00D7 100%. "
                "Target = 0%."],
               ["Verifikasi notifikasi fail-closed",
                "Sifat pemeriksaan notifikasi yang menolak permintaan apabila "
                "tanda tangan tidak sah, tidak tersedia, atau nominal tidak "
                "sesuai; kegagalan tidak pernah dianggap sebagai keberhasilan.",
                "Kode balasan HTTP 401 untuk notifikasi tanpa tanda tangan "
                "sah; nilai status transaksi tetap 'pending'."],
               ["Dua jalur penyelesaian",
                "Tersedianya dua cara untuk mengetahui dan menindaklanjuti "
                "status pembayaran, yaitu notifikasi (webhook) dari Midtrans "
                "dan pemeriksaan status yang dilakukan server sendiri ke "
                "Midtrans Core API.",
                "Kedua jalur menghasilkan status akhir yang sama pada "
                "tabel transactions."],
               ["Registrasi dua fase",
                "Prosedur pendaftaran yang menunda pembuatan baris pengguna "
                "sampai kode OTP terbukti sah, dengan menyimpan sementara "
                "data pendaftaran pada kolom payload (JSONB).",
                "Tidak ada baris pada tabel users sebelum OTP terverifikasi; "
                "jumlah akun belum terverifikasi = 0."],
               ["Tembolok selektif PWA",
                "Kebijakan service worker yang secara sadar melewati "
                "(bypass) tembolok untuk konten berbayar, autentikasi, "
                "langganan, dan transaksi.",
                "Tidak ada entri tembolok untuk /api/books/*/content, "
                "/api/auth/*, /api/subscriptions/*, dan /api/transactions."],
               ["Kesesuaian fungsional",
                "Tingkat kelengkapan dan ketepatan fungsi produk dalam "
                "memenuhi kebutuhan yang ditetapkan, diukur dengan pengujian "
                "black box berbasis ISO/IEC 25010.",
                "Persentase skenario pengujian yang berstatus valid. "
                "Target = 100%."],
               ["Kebergunaan (usability)",
                "Tingkat kemudahan produk untuk dipelajari dan digunakan, "
                "diukur dengan kuesioner System Usability Scale (SUS).",
                "Skor SUS 0\u2013100 yang dikonversi dari 10 butir "
                "pernyataan. Target \u2265 68."]],
              widths=[3.4, 5.6, 5.0], fs=9)
    caption(doc, "Tabel 1.2 Definisi operasional variabel penelitian")

    # ---------- G ----------
    h2(doc, "G. Asumsi dan Keterbatasan Penelitian")
    h3(doc, "1. Asumsi Penelitian")
    bullets(doc, [
        "Midtrans menyediakan layanan Core API QRIS dan mengirimkan "
        "notifikasi dengan tanda tangan SHA-512 sesuai dokumentasinya.",
        "Pengguna memiliki perangkat dengan peramban modern yang mendukung "
        "service worker dan Web App Manifest.",
        "Koneksi internet tersedia pada saat transaksi dan aktivasi "
        "langganan dilakukan; kondisi luring hanya didukung untuk cangkang "
        "aplikasi dan data katalog publik yang telah tersimpan.",
        "Kode OTP dikirimkan melalui surel dan dapat diterima pengguna dalam "
        "waktu kurang dari 10 menit.",
        "Subjek uji coba bersedia memberikan penilaian secara jujur melalui "
        "kuesioner SUS.",
    ], numbering="number")

    h3(doc, "2. Keterbatasan Penelitian")
    body(doc, "Penelitian ini memiliki keterbatasan yang dinyatakan secara "
              "terbuka agar hasilnya dapat ditafsirkan dengan tepat.",
         indent=True, space_after=4)
    bullets(doc, [
        "Pengujian dilakukan pada akun sandbox Midtrans, bukan pada "
        "transaksi uang sungguhan; nominal transaksi yang diuji bersifat "
        "simulasi.",
        "Belum tersedia pengujian otomatis (unit, integrasi, maupun ujung "
        "ke ujung) pada kode sumber produk; pengujian dilakukan secara "
        "manual dengan skenario black box dan skrip permintaan berulang.",
        "Fungsi pencarian buku masih dilakukan di sisi klien atas data yang "
        "sudah diunduh; belum memanfaatkan pencarian teks lengkap "
        "berbasis tsvector dan indeks GIN pada PostgreSQL.",
        "Berkas buku digital masih disimpan sebagai data URL berenkode "
        "base64 pada kolom bertipe TEXT di basis data, sehingga kurang "
        "efisien untuk berkas berukuran besar; pemindahan ke layanan "
        "penyimpanan objek direncanakan sebagai pengembangan lanjut.",
        "Kebijakan keamanan konten (Content Security Policy) belum "
        "diaktifkan pada konfigurasi helmet; aktivasi CSP bersama dengan "
        "penyimpanan token sesi pada kuki HttpOnly direncanakan sebagai "
        "pengembangan lanjut.",
        "Penelitian tidak mencakup penyediaan sistem rekomendasi buku "
        "berbasis kesamaan konten maupun perilaku pengguna.",
        "Lingkungan pengembangan masih bergantung pada sistem operasi "
        "Windows dan belum dikemas dengan wadah (container) atau alur "
        "integrasi berkelanjutan.",
        "Uji coba pengguna dilakukan pada mahasiswa Universitas Wahidiyah "
        "Kediri, sehingga hasil penilaian usability belum dapat digeneralisasi "
        "pada populasi pengguna umum.",
    ], numbering="number")


# ======================== BAB II ========================

def build_bab2(doc):
    bab(doc, "II", "KAJIAN PUSTAKA")

    h2(doc, "A. Kajian Teoretis")

    h3(doc, "1. Rekayasa Perangkat Lunak dan Model Prototyping")
    body(doc, "Rekayasa perangkat lunak adalah disiplin yang menerapkan "
              "prinsip-prinsip rekayasa secara sistematis untuk "
              "mengembangkan, mengoperasikan, dan memelihara perangkat lunak "
              "yang dapat diandalkan dan efisien (Pressman & Maxim, 2021). "
              "Salah satu model proses yang lazim dipakai pada penelitian "
              "rancang bangun adalah model prototyping. Pada model ini, "
              "kebutuhan belum sepenuhnya dapat dirumuskan di awal, sehingga "
              "dibangun purwarupa yang kemudian dievaluasi dan disempurnakan "
              "secara berulang sampai pengguna menyatakan kebutuhannya "
              "terpenuhi. Model ini dipilih karena objek penelitian berupa "
              "platform yang belum pernah ada di lingkungan Universitas "
              "Wahidiyah Kediri, sehingga kebutuhan pengguna perlu "
              "dikristalkan melalui interaksi langsung dengan purwarupa yang "
              "berjalan.")
    body(doc, "Sugiyono (2021) menjelaskan bahwa penelitian dan "
              "pengembangan (Research and Development) adalah metode "
              "penelitian yang digunakan untuk menghasilkan produk tertentu "
              "dan menguji keberhasilan produk tersebut. Tahapan R&D mencakup "
              "pengumpulan informasi, perancangan produk, validasi desain, "
              "revisi desain, uji coba produk, revisi produk, uji coba "
              "pemakaian, dan pembuatan produk akhir. Kerangka tersebut "
              "digunakan sebagai acuan penyusunan prosedur penelitian pada "
              "BAB III.")

    h3(doc, "2. Arsitektur Tiga Lapis dan REST API")
    body(doc, "Arsitektur tiga lapis (three-tier architecture) memisahkan "
              "sistem menjadi lapis presentasi, lapis logika aplikasi, dan "
              "lapis data. Pemisahan ini memudahkan pemeliharaan karena "
              "perubahan pada satu lapis tidak secara langsung mengubah lapis "
              "lainnya. Komunikasi antarlapis dalam penelitian ini "
              "menggunakan gaya REST (Representational State Transfer) dengan "
              "representasi JSON, sebagaimana lazim diterapkan pada layanan "
              "web modern (Muhammad & Paputungan, 2024). Arsitektur sistem "
              "yang dirancang disajikan pada Gambar 2.1.")
    figure(doc, "fig_arsitektur.png",
           "Gambar 2.1 Arsitektur sistem tiga lapis platform Wahidiyah Book",
           width_cm=14.5)

    h3(doc, "3. Single Page Application dengan React")
    body(doc, "Single page application (SPA) adalah aplikasi web yang "
              "memuat satu dokumen HTML dan memperbarui tampilannya secara "
              "dinamis tanpa memuat ulang seluruh halaman. Pendekatan ini "
              "memberikan pengalaman yang mendekati aplikasi asli karena "
              "perpindahan antarhalaman tidak disertai jeda pemuatan. React "
              "adalah pustaka JavaScript untuk membangun antarmuka berbasis "
              "komponen; perbandingan terhadap Angular dan Vue.js menunjukkan "
              "bahwa React unggul pada keluwesan komposisi komponen dan "
              "dukungan ekosistem untuk aplikasi satu halaman (Sekhar "
              "Emmanni, 2023). Penelitian ini menggunakan React 18 dengan "
              "Vite 6 sebagai perakit (bundler).")

    h3(doc, "4. Progressive Web App dan Service Worker")
    body(doc, "Progressive web app (PWA) adalah aplikasi web yang "
              "memanfaatkan service worker dan Web App Manifest sehingga "
              "dapat dipasang di perangkat dan sebagian tetap berfungsi saat "
              "jaringan terputus. Service worker bertindak sebagai proksi "
              "yang berdiri di antara aplikasi dan jaringan, sehingga "
              "kebijakan penyimpanan dapat diatur secara terperinci "
              "(Correia dkk., 2021; Tahir dkk., 2021; Widya & Wakhidah, "
              "2022). Aspek yang jarang dibahas pada penerapan PWA adalah "
              "bahwa kebijakan tembolok yang seragam dapat membocorkan konten "
              "berbayar dan data autentikasi pada perangkat bersama. Oleh "
              "karena itu, penelitian ini merancang kebijakan tembolok "
              "selektif: service worker secara sadar melewati tembolok untuk "
              "konten buku berbayar, endpoint autentikasi, data langganan, "
              "dan data transaksi, sedangkan data katalog publik dilayani "
              "dengan strategi network-first yang berjatuhan ke tembolok.")

    h3(doc, "5. Payment Gateway QRIS dan Midtrans Core API")
    body(doc, "QRIS adalah standardisasi kode QR pembayaran yang ditetapkan "
              "oleh Bank Indonesia sehingga satu kode QR dapat dibayar "
              "melalui aplikasi apa pun yang tergabung dalam sistem tersebut. "
              "Pertumbuhan penggunaan QRIS di Indonesia membuka peluang "
              "sekaligus tantangan bagi sistem transaksi skala kecil "
              "(Rachman dkk., 2024; Gea & Al-Azhar, 2021). Midtrans Core API "
              "menyediakan layanan pembuatan transaksi QRIS (charge), "
              "pemeriksaan status transaksi, dan pengiriman notifikasi "
              "(webhook) ke titik akhir yang ditentukan. Dua sifat penting "
              "dari layanan semacam ini adalah: (a) pengiriman notifikasi "
              "bersifat at-least-once, artinya notifikasi yang sama dapat "
              "dikirim lebih dari satu kali; dan (b) notifikasi dapat hilang "
              "apabila titik akhir penerima tidak tersedia. Kedua sifat ini "
              "menjadi dasar perancangan mekanisme penyelesaian transaksi "
              "pada penelitian ini.")

    h3(doc, "6. Idempotensi dan Pola Compare-and-Swap pada Transaksi "
            "Terdistribusi")
    body(doc, "Suatu operasi disebut idempoten apabila menjalankannya "
              "berkali-kali menghasilkan keadaan akhir yang sama dengan "
              "menjalankannya sekali. Dalam konteks sistem pembayaran "
              "terdistribusi, idempotensi merupakan mekanisme utama untuk "
              "mencegah pemrosesan ganda (duplicate transaction processing) "
              "akibat pengiriman ulang pesan (Journal of Computational "
              "Analysis and Applications, 2025). Salah satu cara praktis "
              "untuk memperoleh idempotensi tanpa menambahkan komponen "
              "pengunci terpisah adalah pola compare-and-swap (CAS) pada "
              "tingkat basis data: pembaruan dijalankan dengan syarat bahwa "
              "keadaan baris masih sama dengan keadaan yang diharapkan, "
              "kemudian jumlah baris yang terpengaruh diperiksa. Apabila "
              "tidak ada baris yang terpengaruh, berarti pekerja lain sudah "
              "mengklaim perubahan tersebut, sehingga pekerja saat ini "
              "cukup menghentikan prosesnya.")
    body(doc, "Dalam penelitian ini, CAS diwujudkan dengan pernyataan "
              "UPDATE yang menyaratkan status masih 'pending', lalu memeriksa "
              "bahwa jumlah baris yang dikembalikan sama dengan satu. Hanya "
              "pemenang klaim yang melanjutkan ke tahap aktivasi langganan. "
              "Apabila aktivasi gagal, dilakukan transaksi kompensasi yang "
              "mengembalikan status menjadi 'pending' dan mengosongkan waktu "
              "verifikasi, sehingga percobaan ulang berikutnya dapat "
              "memproses transaksi yang sama. Alur tersebut disajikan pada "
              "Gambar 2.2.")
    figure(doc, "fig_idempoten.png",
           "Gambar 2.2 Alur penyelesaian transaksi QRIS yang idempoten "
           "dengan compare-and-swap dan kompensasi", width_cm=13.5)

    h3(doc, "7. Dua Jalur Penyelesaian Transaksi")
    body(doc, "Ketergantungan tunggal pada notifikasi (webhook) membuat "
              "sistem rentan: apabila notifikasi hilang, status transaksi "
              "tidak pernah berubah. Penelitian ini merancang dua jalur yang "
              "bermuara pada satu fungsi penyelesaian yang sama. Jalur A "
              "adalah notifikasi yang dikirim Midtrans ke titik akhir server. "
              "Jalur B adalah titik akhir pemeriksaan status yang diakses "
              "klien, di mana server sendiri memanggil Midtrans Core API "
              "untuk menanyakan status transaksi, lalu menjalankan "
              "penyelesaian yang sama. Karena fungsi penyelesaian bersifat "
              "idempoten, kedua jalur dapat berjalan berbarengan tanpa "
              "menghasilkan aktivasi ganda. Jalur B sekaligus memecahkan "
              "masalah pengembangan lokal: server pada localhost tidak "
              "memiliki URL publik, namun pemeriksaan status tetap dapat "
              "dilakukan karena yang memanggil Midtrans adalah server, bukan "
              "sebaliknya. Rancangan ini disajikan pada Gambar 2.3.")
    figure(doc, "fig_dualpath.png",
           "Gambar 2.3 Dua jalur penyelesaian transaksi: webhook dan polling "
           "otoritatif server", width_cm=14.5)

    h3(doc, "8. Verifikasi Notifikasi yang Fail-Closed")
    body(doc, "Verifikasi notifikasi pembayaran harus dirancang dengan "
              "prinsip fail-closed, yaitu kegagalan verifikasi menghasilkan "
              "penolakan, bukan penerimaan. Penelitian ini menghitung nilai "
              "tanda tangan sebagai SHA-512 atas gabungan order_id, "
              "status_code, gross_amount, dan kunci server, kemudian "
              "membandingkannya dengan nilai yang dikirim menggunakan fungsi "
              "perbandingan waktu konstan. Pemeriksaan panjang dilakukan "
              "sebelum perbandingan agar fungsi tidak melemparkan kesalahan "
              "dan agar tidak terjadi kebocoran informasi melalui perbedaan "
              "waktu. Selain tanda tangan, dilakukan rekonsiliasi nominal "
              "untuk memastikan jumlah yang dilaporkan sama dengan jumlah "
              "yang tersimpan pada basis data, serta pembatasan host untuk "
              "pengambilan berkas gambar kode QR guna mencegah penyalahgunaan "
              "pengambilan sumber daya.")

    h3(doc, "9. Autentikasi: JWT, bcrypt, dan Pendaftaran Dua Fase")
    body(doc, "JSON Web Token (JWT) adalah standar terbuka untuk "
              "merepresentasikan klaim secara ringkas dan aman antarpihak, "
              "yang lazim digunakan sebagai token sesi pada layanan web "
              "(Pudoli dkk., 2023; Gustiegan & Painem, 2022; Rana, 2023). Agar "
              "token tidak dapat dipalsukan, token ditandatangani dengan "
              "kunci rahasia di sisi server. Kata sandi tidak disimpan dalam "
              "bentuk asli, melainkan dalam bentuk nilai hash yang dihasilkan "
              "dengan bcrypt menggunakan 12 putaran, sehingga biaya komputasi "
              "untuk percobaan paksa menjadi tinggi.")
    body(doc, "Penelitian ini menerapkan pendaftaran dua fase yang berbeda "
              "dari pola lazim. Pada fase pertama, server tidak membuat "
              "baris pada tabel users sama sekali. Server memvalidasi "
              "masukan, menghitung hash kata sandi, lalu menyimpan "
              "sementara nama, hash kata sandi, dan tanggal lahir pada kolom "
              "bertipe JSONB di tabel verifikasi surel bersama kode OTP "
              "bermasa berlaku 10 menit. Baris pengguna baru dibuat pada "
              "fase kedua, yaitu setelah kode OTP terbukti sah. Dengan "
              "demikian basis data tidak pernah memuat akun yang belum "
              "terverifikasi, dan pelanggaran keterbatasan unik yang mungkin "
              "terjadi pada saat penyisipan dapat ditangani secara ramah. "
              "Rancangan ini disajikan pada Gambar 2.4.")
    figure(doc, "fig_registrasi.png",
           "Gambar 2.4 Registrasi dua fase berbasis OTP", width_cm=13.5)

    h3(doc, "10. Pembatasan Akses pada Tingkat Baris (Row Level Security)")
    body(doc, "Row Level Security (RLS) pada PostgreSQL memungkinkan "
              "kebijakan akses dinyatakan sebagai predikat SQL yang "
              "dievaluasi untuk setiap baris, sehingga pembatasan hak akses "
              "ditegakkan oleh mesin basis data, bukan hanya oleh kode "
              "aplikasi. Hal ini penting sebagai pertahanan berlapis: apabila "
              "terdapat kelemahan pada kode aplikasi, basis data tetap "
              "menegakkan aturannya. Penelitian ini memanfaatkan RLS bersama "
              "dengan sebuah trigger yang berjalan sebelum penyisipan atau "
              "pembaruan pada tabel transaksi QRIS untuk menjaga agar "
              "perpindahan keadaan selalu sah.")

    h3(doc, "11. Kualitas Perangkat Lunak ISO/IEC 25010 dan Pengujian "
            "Black Box")
    body(doc, "ISO/IEC 25010 mendefinisikan model kualitas produk "
              "perangkat lunak yang terdiri atas kesesuaian fungsional, "
              "keandalan, efisiensi kinerja, kegunaan, keamanan, "
              "kompatibilitas, kemudahan pemeliharaan, dan portabilitas. "
              "Pengujian black box adalah teknik pengujian yang menilai "
              "perangkat lunak dari sisi perilaku eksternal tanpa melihat "
              "struktur kode; teknik ini lazim dipadukan dengan standar "
              "ISO/IEC 25010 untuk mengukur kualitas situs dan aplikasi web "
              "(Hasibuan dkk., 2025; Minarni & Sigit, 2023; Alexandro & "
              "Parhusip, 2025). Penelitian ini menyusun skenario pengujian "
              "yang memetakan setiap fungsi utama ke karakteristik ISO/IEC "
              "25010 yang relevan.")

    h3(doc, "12. System Usability Scale")
    body(doc, "System Usability Scale (SUS) adalah kuesioner sepuluh butir "
              "dengan skala Likert lima poin yang menghasilkan skor 0\u2013"
              "100 untuk mengukur persepsi kebergunaan suatu sistem. SUS "
              "banyak digunakan karena ringkas, tidak bergantung pada "
              "perangkat, dan telah teruji secara statistik (Putri, 2022; "
              "Kusumaningrum dkk., 2023; Alfilaeni dkk., 2024). Skor SUS "
              "bukan persentase, melainkan posisi pada skala; nilai 68 "
              "umumnya dijadikan ambang keberterimaan (acceptable).")

    # ---------- B ----------
    h2(doc, "B. Hasil Penelitian yang Relevan")
    body(doc, "Penelitian ini diposisikan terhadap sejumlah penelitian "
              "terdahulu yang sebidang. Perbandingan berikut menunjukkan "
              "persamaan dan perbedaan, sekaligus menegaskan kebaruan "
              "(novelty) penelitian ini sebagaimana disajikan pada Tabel 2.1.",
         indent=True, space_after=8)
    add_table(doc,
              ["Peneliti (Tahun)", "Fokus & Metode", "Hasil",
               "Posisi Penelitian Ini"],
              [["Mulyani & Hati (2021)",
                "Rancang bangun sistem informasi perpustakaan digital "
                "berbasis web; model waterfall.",
                "Katalog perpustakaan dapat diakses daring dan mempercepat "
                "penelusuran koleksi.",
                "Sejalan pada modul katalog, namun penelitian ini menambahkan "
                "ruang baca digital berlangganan dan transaksi QRIS."],
               ["Saekup (2022)",
                "Rancang bangun sistem informasi perpustakaan berbasis web "
                "di sekolah menengah kejuruan.",
                "Proses peminjaman dan pengembalian tercatat otomatis.",
                "Menambahkan pembatasan akses konten berbayar dan "
                "pengamanan tembolok PWA."],
               ["Wijana & Prasetio (2023)",
                "Rancang bangun sistem informasi perpustakaan berbasis "
                "situs web.",
                "Sirkulasi dan penelusuran buku terdigitalisasi.",
                "Menambahkan integrasi pembayaran dan analisis keandalan "
                "transaksi."],
               ["Muttaqin dkk. (2022)",
                "Rancang bangun aplikasi e-commerce dengan payment dan "
                "shipping gateway.",
                "Pembayaran daring terintegrasi dengan toko.",
                "Tidak menguji idempotensi; penelitian ini menjadikannya "
                "fokus utama."],
               ["Nariswari dkk. (2023)",
                "Rancang bangun sistem e-tiket berbasis payment gateway.",
                "Pemesanan dan pembayaran tiket terotomatisasi.",
                "Tidak membahas notifikasi ganda; penelitian ini merancang "
                "CAS dan kompensasi."],
               ["Rezkiandra (2024)",
                "Rancang bangun aplikasi warung digital berbasis web "
                "menggunakan payment gateway.",
                "Transaksi penjualan tercatat dan dapat dipantau.",
                "Tidak membahas pengujian lokal tanpa URL publik; penelitian "
                "ini menyediakan jalur polling server."],
               ["Richo dkk. (2021)",
                "Rancang bangun aplikasi point of sales dengan payment "
                "gateway terintegrasi.",
                "Pembayaran kasir terintegrasi dengan gateway.",
                "Tidak menerapkan verifikasi fail-closed; penelitian ini "
                "menerapkan SHA-512 dan timingSafeEqual."],
               ["Hambali dkk. (2024)",
                "Perancangan dan evaluasi sistem manajemen keanggotaan "
                "berbasis web.",
                "Pengelolaan masa berlaku keanggotaan terotomatisasi.",
                "Sejalan pada modul langganan; penelitian ini menambahkan "
                "pencegahan aktivasi ganda."]],
              widths=[2.8, 3.6, 3.6, 4.0], fs=9)
    caption(doc, "Tabel 2.1 Hasil penelitian terdahulu yang relevan")

    body(doc, "Berdasarkan Tabel 2.1 tampak bahwa penelitian rancang bangun "
              "perpustakaan digital umumnya berhenti pada digitalisasi "
              "katalog dan sirkulasi, sedangkan penelitian integrasi payment "
              "gateway umumnya berhenti pada pernyataan bahwa pembayaran "
              "berhasil dilakukan. Kebaruan penelitian ini terletak pada "
              "perancangan dan pengujian mekanisme yang menjaga konsistensi "
              "status transaksi pembayaran, yaitu penyelesaian idempoten "
              "berbasis compare-and-swap dengan transaksi kompensasi, "
              "verifikasi notifikasi yang fail-closed, dan dua jalur "
              "penyelesaian yang memungkinkan pengujian tanpa URL publik.")

    # ---------- C ----------
    h2(doc, "C. Kerangka Konseptual")
    body(doc, "Kerangka konseptual menggambarkan hubungan antara komponen "
              "masukan, proses, dan luaran dari sistem yang dirancang. "
              "Masukan penelitian berupa pengguna dengan peran yang berbeda, "
              "data katalog buku, paket langganan, dan notifikasi pembayaran. "
              "Proses mencakup autentikasi dua fase, penyajian katalog dan "
              "ruang baca digital dengan pembatasan hak akses, transaksi QRIS "
              "yang diselesaikan secara idempoten, kebijakan tembolok PWA "
              "yang selektif, serta dasbor pengelola. Luaran yang diharapkan "
              "berupa akun terverifikasi, langganan aktif tanpa penggandaan, "
              "konten berbayar yang terlindungi, log aktivitas pengelola, "
              "serta ukuran kualitas fungsional dan kebergunaan produk. "
              "Kerangka konseptual disajikan pada Gambar 2.5.")
    figure(doc, "fig_konseptual.png",
           "Gambar 2.5 Kerangka konseptual penelitian", width_cm=15.0)

    # ---------- D ----------
    h2(doc, "D. Kerangka Model Penelitian")
    body(doc, "Kerangka model penelitian berikut merangkum keterkaitan "
              "antara latar belakang masalah, rumusan masalah, variabel, "
              "metode, langkah pelaksanaan, dan instrumen yang digunakan "
              "dalam penelitian ini. Rangkuman tersebut disajikan pada Tabel "
              "2.2.", indent=True, space_after=8)

    add_table(doc,
              ["Latar Belakang", "Masalah", "Definisi Operasional Variabel",
               "Metode", "Langkah-langkah", "Instrumen"],
              [["1. Pengelolaan katalog dan ruang baca di lingkungan "
                "Universitas Wahidiyah Kediri masih terpisah-pisah dan belum "
                "terintegrasi dengan pembayaran digital.\n\n"
                "2. Pengguna kesulitan menemukan bahan bacaan dan tidak "
                "tersedia ruang baca digital berlangganan.\n\n"
                "3. Notifikasi payment gateway bersifat at-least-once dan "
                "dapat hilang, sehingga berisiko menimbulkan aktivasi "
                "langganan ganda atau pembayaran yang tidak tercatat.\n\n"
                "4. Pengembangan lokal tanpa URL publik menyulitkan "
                "pengujian alur pembayaran.",

                "1. Bagaimana merancang dan membangun platform Wahidiyah "
                "Book yang mengintegrasikan katalog, ruang baca digital, dan "
                "langganan berbayar?\n"
                "2. Bagaimana merancang penyelesaian transaksi QRIS yang "
                "idempoten?\n"
                "3. Bagaimana merancang verifikasi notifikasi yang "
                "fail-closed?\n"
                "4. Bagaimana menyediakan dua jalur penyelesaian transaksi?\n"
                "5. Bagaimana mencegah terbentuknya akun belum "
                "terverifikasi?\n"
                "6. Bagaimana mencegah tersimpannya konten berbayar di "
                "tembolok PWA?\n"
                "7. Bagaimana kualitas fungsional dan kebergunaan produk?",

                "1. Penyelesaian transaksi idempoten: keadaan akhir sama "
                "berapa pun kali operasi dijalankan.\n"
                "2. Rasio aktivasi ganda: (aktivasi ganda / seluruh "
                "order_id) \u00D7 100%; target 0%.\n"
                "3. Verifikasi fail-closed: notifikasi tanpa tanda tangan sah "
                "ditolak dengan HTTP 401.\n"
                "4. Dua jalur penyelesaian: webhook dan pemeriksaan status "
                "otoritatif server.\n"
                "5. Registrasi dua fase: baris users dibuat hanya setelah OTP "
                "sah.\n"
                "6. Tembolok selektif PWA: konten berbayar dan data sensitif "
                "dilewati dari tembolok.\n"
                "7. Kesesuaian fungsional (ISO/IEC 25010) dan kebergunaan "
                "(SUS).",

                "1. Jenis: Research and Development (rancang bangun) dengan "
                "model prototyping.\n"
                "2. Tempat: lingkungan Universitas Wahidiyah Kediri.\n"
                "3. Pendekatan: arsitektur tiga lapis, REST + JSON, SPA "
                "dengan React, basis data PostgreSQL.\n"
                "4. Subjek uji coba: pengelola, mahasiswa pengguna, dan "
                "penilai ahli.\n"
                "5. Teknik pengumpulan data: skenario black box, log server, "
                "kuesioner SUS, wawancara.",

                "1. Studi pendahuluan: observasi, wawancara, studi "
                "pustaka.\n"
                "2. Perancangan: use case, ERD, arsitektur, kontrak API, "
                "antarmuka.\n"
                "3. Implementasi: klien React, server Express, skema basis "
                "data.\n"
                "4. Pengujian internal: black box berbasis ISO/IEC 25010 dan "
                "eksperimen idempotensi.\n"
                "5. Uji coba pengguna: skenario tugas dan kuesioner SUS.\n"
                "6. Evaluasi dan penyempurnaan iteratif sampai spesifikasi "
                "terpenuhi.",

                "1. Lembar skenario pengujian black box.\n"
                "2. Skrip permintaan berulang untuk uji idempotensi.\n"
                "3. Log server dan catatan basis data.\n"
                "4. Kuesioner System Usability Scale (10 butir).\n"
                "5. Pedoman wawancara pengelola.\n"
                "6. Dokumentasi kode sumber dan skema basis data."]],
              widths=[2.3, 2.4, 2.5, 2.4, 2.3, 2.1], fs=8.5)
    caption(doc, "Tabel 2.2 Kerangka model penelitian rancang bangun "
                 "Wahidiyah Book")

    # ---------- E ----------
    h2(doc, "E. Hipotesis")
    body(doc, "Mengingat penelitian ini berjenis pengembangan (rancang "
              "bangun), hipotesis yang dirumuskan bukan hipotesis hubungan "
              "antarvariabel, melainkan pernyataan keberhasilan produk yang "
              "akan diuji pada tahap uji coba. Adapun pernyataan keberhasilan "
              "tersebut adalah sebagai berikut.", indent=True, space_after=4)
    bullets(doc, [
        "H\u2081: Platform Wahidiyah Book yang dirancang memenuhi seluruh "
        "skenario pengujian black box pada fungsi utama (kesesuaian "
        "fungsional 100%).",
        "H\u2082: Mekanisme penyelesaian transaksi berbasis compare-and-swap "
        "menghasilkan rasio aktivasi ganda sebesar 0% pada sepuluh "
        "pengiriman ulang notifikasi untuk satu order_id yang sama.",
        "H\u2083: Verifikasi notifikasi yang dirancang menolak seluruh "
        "notifikasi tanpa tanda tangan sah (fail-closed).",
        "H\u2084: Kedua jalur penyelesaian, yaitu webhook dan pemeriksaan "
        "status otoritatif server, menghasilkan keadaan akhir transaksi yang "
        "sama.",
        "H\u2085: Platform yang dihasilkan memperoleh skor System Usability "
        "Scale \u2265 68 dengan tingkat penyelesaian tugas \u2265 90%.",
    ], numbering="number")


# ======================== BAB III ========================

def build_bab3(doc):
    bab(doc, "III", "METODE PENELITIAN")

    h2(doc, "A. Jenis Penelitian")
    body(doc, "Jenis penelitian yang digunakan adalah penelitian dan "
              "pengembangan (Research and Development / R&D) dengan "
              "pendekatan rancang bangun. Penelitian R&D dipilih karena "
              "tujuan penelitian adalah menghasilkan produk berupa perangkat "
              "lunak beserta bukti bahwa produk tersebut memenuhi "
              "spesifikasi yang ditetapkan (Sugiyono, 2021). Model proses "
              "yang digunakan adalah model prototyping (Pressman & Maxim, "
              "2021), yaitu model yang membangun purwarupa secara bertahap, "
              "mengevaluasinya bersama pengguna, lalu menyempurnakannya "
              "dalam siklus berulang.")
    body(doc, "Pemilihan model prototyping didasarkan pada tiga "
              "pertimbangan. Pertama, platform yang dibangun belum memiliki "
              "padanan di lingkungan Universitas Wahidiyah Kediri, sehingga "
              "kebutuhan tidak dapat dirumuskan tuntas di awal dan perlu "
              "dikristalkan melalui purwarupa yang berjalan. Kedua, "
              "kebutuhan pengguna bersifat dinamis dan sangat dipengaruhi "
              "oleh umpan balik terhadap antarmuka yang nyata. Ketiga, model "
              "ini memungkinkan pengujian komponen berisiko tinggi, yaitu "
              "alur pembayaran, dilakukan lebih awal sehingga kelemahan "
              "desain dapat ditemukan sebelum seluruh fitur selesai.")
    body(doc, "Tahapan prosedur pengembangan yang akan ditempuh disajikan "
              "pada Gambar 3.1.")
    figure(doc, "fig_prosedur.png",
           "Gambar 3.1 Tahapan prosedur pengembangan model prototyping",
           width_cm=15.0)

    h2(doc, "B. Prosedur Penelitian")
    body(doc, "Prosedur penelitian mengikuti tahapan R&D yang disesuaikan "
              "dengan kebutuhan rekayasa perangkat lunak. Setiap tahap "
              "dijabarkan berikut ini.", indent=True, space_after=6)

    h3(doc, "1. Tahap Studi Pendahuluan dan Analisis Kebutuhan")
    bullets(doc, [
        "Observasi terhadap kondisi pengelolaan bahan bacaan dan layanan "
        "perpustakaan di lingkungan Universitas Wahidiyah Kediri.",
        "Wawancara dengan pengelola untuk menggali alur bisnis, peran "
        "pengguna, dan kendala yang dihadapi.",
        "Studi pustaka terhadap teori rekayasa perangkat lunak, arsitektur "
        "web, pembayaran QRIS, idempotensi, dan pengukuran kualitas produk.",
        "Perumusan kebutuhan fungsional dan non-fungsional yang dituangkan "
        "ke dalam spesifikasi produk sebagaimana Tabel 1.1.",
    ], numbering="number")

    h3(doc, "2. Tahap Perancangan")
    bullets(doc, [
        "Perancangan use case untuk aktor tamu, pengguna, dan pengelola.",
        "Perancangan skema basis data (ERD) yang memuat tabel users, books, "
        "events, subscriptions, transactions, password_resets, "
        "email_verifications, qris_transactions, dan carousel_slides.",
        "Perancangan arsitektur tiga lapis dan kontrak antarmuka pemrograman "
        "aplikasi (API) untuk setiap endpoint yang disediakan.",
        "Perancangan antarmuka pengguna (wireframe) untuk halaman katalog, "
        "detail buku, ruang baca, pembayaran, dasbor pengelola, dan profil.",
        "Perancangan mekanisme penyelesaian transaksi, meliputi alur "
        "compare-and-swap, transaksi kompensasi, verifikasi notifikasi, dan "
        "dua jalur penyelesaian.",
    ], numbering="number")

    h3(doc, "3. Tahap Implementasi")
    body(doc, "Implementasi dilakukan dengan susunan teknologi sebagai "
              "berikut.", indent=True, space_after=4)
    add_table(doc,
              ["Lapis", "Teknologi", "Keterangan"],
              [["Presentasi", "React 18, Vite 6, Tailwind CSS, "
                              "react-router-dom, Recharts, qrcode.react",
                "Aplikasi satu halaman yang dapat dipasang sebagai PWA."],
               ["Logika aplikasi", "Node.js, Express 4 (ESM), helmet, "
                                   "express-rate-limit, zod, jsonwebtoken, "
                                   "bcryptjs, multer, sharp, dayjs",
                "Satu peladen untuk API sekaligus penyedia berkas hasil "
                "bangun."],
               ["Basis data", "PostgreSQL (Supabase), Row Level Security, "
                              "trigger fn_handle_qris_payment()",
                "Invarians keadaan ditegakkan oleh mesin basis data."],
               ["Pembayaran", "Midtrans Core API (QRIS)",
                "Pembuatan kode QR, pemeriksaan status, dan notifikasi."],
               ["Surel", "Gmail API dengan pembangun pesan MIME RFC 2822",
                "Pengiriman kode OTP dan pemberitahuan."]],
              widths=[2.6, 6.0, 5.4], fs=9.5)
    caption(doc, "Tabel 3.1 Susunan teknologi implementasi", before=6)

    body(doc, "Tiga bagian implementasi yang menjadi inti kontribusi "
              "penelitian diuraikan berikut ini.", indent=True, space_after=4)

    h3(doc, "a. Penyelesaian transaksi idempoten")
    body(doc, "Fungsi penyelesaian transaksi dirancang sedemikian rupa "
              "sehingga aman dipanggil dari jalur mana pun dan berapa pun "
              "kali. Inti mekanismenya adalah pembaruan bersyarat yang hanya "
              "berhasil apabila keadaan transaksi masih 'pending', lalu "
              "pemeriksaan jumlah baris yang terpengaruh. Hanya pemenang "
              "klaim yang menjalankan aktivasi langganan; apabila aktivasi "
              "gagal, status dikembalikan ke 'pending' agar percobaan ulang "
              "dapat memprosesnya kembali. Cuplikan inti mekanisme tersebut "
              "diberikan pada Lampiran 3.")

    h3(doc, "b. Verifikasi notifikasi yang fail-closed")
    body(doc, "Setiap notifikasi yang masuk diverifikasi sebelum ditindak "
              "lanjuti. Nilai yang diharapkan dihitung sebagai SHA-512 atas "
              "gabungan order_id, status_code, gross_amount, dan kunci "
              "server; hasilnya dibandingkan dengan nilai yang dikirim "
              "menggunakan perbandingan waktu konstan setelah pemeriksaan "
              "panjang. Apabila kunci server belum tersedia, fungsi langsung "
              "menyatakan tidak valid alih-alih mengizinkan permintaan lewat. "
              "Setelah tanda tangan sah, dilakukan rekonsiliasi antara nominal "
              "yang dilaporkan dengan nominal yang tersimpan.")

    h3(doc, "c. Registrasi dua fase berbasis OTP")
    body(doc, "Permintaan pendaftaran tidak langsung membuat baris pengguna. "
              "Server memvalidasi masukan dengan zod, menghitung hash kata "
              "sandi dengan bcrypt pada 12 putaran, lalu menyimpan nama, "
              "hash kata sandi, dan tanggal lahir pada kolom JSONB tabel "
              "verifikasi surel bersama kode OTP bermasa berlaku 10 menit "
              "dan pembatasan lima percobaan. Baris pengguna dibuat hanya "
              "setelah kode OTP terbukti sah, dengan penanganan khusus "
              "apabila terjadi pelanggaran keterbatasan unik.")

    h3(doc, "4. Tahap Pengujian Internal")
    bullets(doc, [
        "Pengujian black box terhadap seluruh fungsi utama dengan teknik "
        "partisi ekuivalensi dan analisis nilai batas.",
        "Eksperimen idempotensi: mengirimkan notifikasi yang sama sebanyak "
        "sepuluh kali untuk satu order_id, serta menjalankan dua jalur "
        "penyelesaian secara berbarengan, lalu memeriksa jumlah aktivasi "
        "langganan yang tercatat.",
        "Pengujian keamanan dasar: mengirimkan notifikasi tanpa tanda "
        "tangan sah, mengakses konten berbayar tanpa langganan, dan "
        "memeriksa isi tembolok service worker.",
        "Pencatatan temuan dan perbaikan, diikuti pengujian ulang sampai "
        "seluruh skenario berstatus valid.",
    ], numbering="number")

    h3(doc, "5. Tahap Uji Coba Pengguna")
    body(doc, "Uji coba pengguna dilakukan dengan memberikan seperangkat "
              "tugas nyata kepada responden, misalnya mendaftar akun, "
              "mencari buku, membaca buku terbuka, melakukan pembayaran "
              "langganan, dan mengelola konten pada dasbor pengelola. "
              "Keberhasilan penyelesaian tugas dicatat, kemudian responden "
              "mengisi kuesioner System Usability Scale.")

    h3(doc, "6. Tahap Evaluasi dan Penyempurnaan Iteratif")
    body(doc, "Hasil pengujian internal dan uji coba pengguna dievaluasi "
              "terhadap spesifikasi produk pada Tabel 1.1. Temuan yang "
              "berdampak pada fungsi, keandalan, keamanan, atau kebergunaan "
              "diperbaiki, kemudian produk dibangun ulang dan diuji kembali. "
              "Siklus ini diulang sampai seluruh target spesifikasi "
              "terpenuhi, setelah itu produk dirilis sebagai hasil akhir "
              "penelitian.")

    h2(doc, "C. Desain Uji Coba")
    body(doc, "Desain uji coba yang digunakan adalah desain pasca-uji "
              "satu kali (one-shot post-test only), yaitu pengukuran "
              "dilakukan setelah produk selesai dibangun dan responden "
              "menggunakannya. Desain ini dipilih karena produk yang diuji "
              "bersifat tunggal dan tidak dibandingkan dengan produk "
              "lainnya. Uji coba mencakup tiga jenis pengujian.")
    bullets(doc, [
        "Pengujian fungsional (black box): menguji kesesuaian keluaran "
        "terhadap masukan pada setiap fungsi utama, dengan pemetaan ke "
        "karakteristik ISO/IEC 25010. Contoh skenario disajikan pada Tabel "
        "3.2.",
        "Eksperimen keandalan (idempotensi): menguji bahwa pengulangan "
        "permintaan penyelesaian tidak menghasilkan aktivasi ganda.",
        "Pengujian kebergunaan (usability): mengukur persepsi responden "
        "dengan kuesioner System Usability Scale dan mencatat tingkat "
        "penyelesaian tugas.",
    ], numbering="number")

    add_table(doc,
              ["No.", "Fungsi yang Diuji", "Skenario", "Keluaran yang "
               "Diharapkan", "Karakteristik ISO/IEC 25010"],
              [["1.", "Pendaftaran akun",
                "Mendaftar dengan surel dan kata sandi yang valid, lalu "
                "memasukkan kode OTP yang benar.",
                "Kode OTP diterima; akun tercipta setelah OTP benar; tidak "
                "ada akun sebelum verifikasi.", "Kesesuaian fungsional, "
                "keamanan"],
               ["2.", "Pendaftaran akun",
                "Memasukkan kode OTP yang salah sebanyak lima kali.",
                "Percobaan dikunci; permintaan berikutnya ditolak.",
                "Keamanan"],
               ["3.", "Masuk (login)",
                "Masuk dengan kata sandi salah.",
                "Pesan kesalahan seragam tanpa memberitahukan apakah surel "
                "terdaftar.", "Keamanan"],
               ["4.", "Pencarian katalog",
                "Mencari kata kunci yang ada dan yang tidak ada.",
                "Daftar buku yang relevan; hasil kosong disertai pesan yang "
                "jelas.", "Kesesuaian fungsional, kebergunaan"],
               ["5.", "Ruang baca digital",
                "Membuka buku berbayar tanpa langganan aktif.",
                "Akses ditolak dan pengguna diarahkan ke halaman "
                "berlangganan.", "Keamanan, kesesuaian fungsional"],
               ["6.", "Transaksi QRIS",
                "Membuat transaksi lalu menyelesaikan pembayaran.",
                "Kode QR tampil; status berubah menjadi berhasil; langganan "
                "aktif.", "Kesesuaian fungsional"],
               ["7.", "Transaksi QRIS",
                "Mengirim notifikasi penyelesaian yang sama sebanyak sepuluh "
                "kali.",
                "Tepat satu aktivasi langganan; sembilan sisanya dinyatakan "
                "sudah diproses.", "Keandalan"],
               ["8.", "Verifikasi notifikasi",
                "Mengirim notifikasi tanpa tanda tangan sah.",
                "Ditolak dengan kode 401; status transaksi tetap 'pending'.",
                "Keamanan"],
               ["9.", "Tembolok PWA",
                "Mengakses konten berbayar lalu memeriksa isi tembolok "
                "service worker.",
                "Tidak ada entri tembolok untuk konten berbayar, "
                "autentikasi, langganan, dan transaksi.", "Keamanan"],
               ["10.", "Dasbor pengelola",
                "Menambah, mengubah, dan menghapus data buku.",
                "Data tersimpan dan tampil konsisten; log aktivitas "
                "tercatat.", "Kesesuaian fungsional, kemudahan pemeliharaan"]],
              widths=[1.0, 2.6, 4.2, 3.9, 2.3], fs=8.5)
    caption(doc, "Tabel 3.2 Contoh skenario pengujian black box")

    h2(doc, "D. Subjek Uji Coba")
    body(doc, "Subjek uji coba dalam penelitian ini terdiri atas tiga "
              "kelompok.", indent=True, space_after=4)
    add_table(doc,
              ["Kelompok", "Jumlah", "Kriteria Inklusi", "Peran dalam Uji "
               "Coba"],
              [["Pengelola (admin)", "1\u20132 orang",
                "Dosen atau staf yang bertanggung jawab atas pengelolaan "
                "konten pada platform.",
                "Menguji dasbor pengelola dan memberikan masukan alur "
                "bisnis."],
               ["Pengguna (mahasiswa)", "\u2265 30 orang",
                "Mahasiswa aktif Universitas Wahidiyah Kediri yang pernah "
                "mengakses aplikasi web dan memiliki aplikasi pembayaran "
                "ber-QRIS.",
                "Melakukan skenario tugas dan mengisi kuesioner SUS."],
               ["Penilai ahli", "2\u20133 orang",
                "Dosen atau praktisi dengan latar belakang rekayasa "
                "perangkat lunak.",
                "Menilai kelayakan desain, keamanan, dan kesesuaian "
                "terhadap ISO/IEC 25010."]],
              widths=[3.2, 2.0, 5.0, 4.8], fs=9.5)
    caption(doc, "Tabel 3.3 Subjek uji coba penelitian")

    h2(doc, "E. Jenis Data dan Sumber Data")
    bullets(doc, [
        "Data primer: hasil pengujian black box, log server dan catatan "
        "basis data pada eksperimen idempotensi, tingkat penyelesaian tugas "
        "responden, skor kuesioner SUS, dan hasil wawancara dengan pengelola "
        "serta penilai ahli.",
        "Data sekunder: dokumentasi Midtrans Core API, dokumentasi React, "
        "Vite, dan PostgreSQL, standar ISO/IEC 25010, serta artikel jurnal "
        "dan buku yang tercantum pada daftar pustaka.",
    ], numbering="number")

    h2(doc, "F. Instrumen Pengumpulan Data")
    body(doc, "Instrumen yang digunakan untuk mengumpulkan data dalam "
              "penelitian ini dirangkum pada Tabel 3.4.",
         indent=True, space_after=8)
    add_table(doc,
              ["Instrumen", "Bentuk", "Digunakan untuk"],
              [["Lembar skenario pengujian black box",
                "Daftar berisi nomor, fungsi, langkah, keluaran yang "
                "diharapkan, dan kolom status valid/tidak valid",
                "Mengukur kesesuaian fungsional produk."],
               ["Skrip permintaan berulang",
                "Skrip yang mengirim permintaan penyelesaian yang sama "
                "berulang kali ke titik akhir notifikasi",
                "Mengukur idempotensi penyelesaian transaksi."],
               ["Log server dan catatan basis data",
                "Berkas log aplikasi dan hasil kueri pada tabel "
                "transactions dan subscriptions",
                "Memverifikasi jumlah aktivasi dan perpindahan keadaan "
                "transaksi."],
               ["Kuesioner System Usability Scale",
                "Sepuluh butir pernyataan dengan skala Likert lima poin "
                "(sangat tidak setuju sampai sangat setuju)",
                "Mengukur tingkat kebergunaan produk."],
               ["Pedoman wawancara",
                "Daftar pertanyaan terbuka untuk pengelola dan penilai ahli",
                "Menggali kebutuhan, kendala, dan saran perbaikan."],
               ["Lembar observasi tugas",
                "Catatan keberhasilan, waktu penyelesaian, dan kendala "
                "responden",
                "Mengukur tingkat penyelesaian tugas dan efisiensi "
                "penggunaan."]],
              widths=[3.4, 5.6, 5.0], fs=9.5)
    caption(doc, "Tabel 3.4 Instrumen pengumpulan data")

    h2(doc, "G. Teknik Analisis Data")
    bullets(doc, [
        "Analisis kesesuaian fungsional: hasil pengujian black box "
        "dinyatakan sebagai persentase skenario yang berstatus valid "
        "terhadap seluruh skenario yang diuji, lalu ditafsirkan terhadap "
        "karakteristik ISO/IEC 25010 yang bersesuaian.",
        "Analisis keandalan (idempotensi): dihitung rasio aktivasi ganda "
        "sebagai jumlah order_id yang menghasilkan lebih dari satu aktivasi "
        "dibagi seluruh order_id yang diuji, dikalikan 100 persen. "
        "Keberhasilan dinyatakan apabila rasio sama dengan 0 persen. Selain "
        "itu dicatat waktu yang dibutuhkan sampai status transaksi menjadi "
        "berhasil.",
        "Analisis keamanan: persentase notifikasi tanpa tanda tangan sah "
        "yang berhasil ditolak, serta pemeriksaan bahwa tidak ada entri "
        "tembolok untuk konten berbayar dan data sensitif.",
        "Analisis kebergunaan: skor SUS dihitung dengan mengurangi satu "
        "dari nilai butir bernomor ganjil dan mengurangi nilai butir "
        "bernomor genap dari lima, menjumlahkan seluruh hasilnya, lalu "
        "mengalikannya dengan 2,5 sehingga diperoleh skor 0\u2013100. Skor "
        "rata-rata kemudian ditafsirkan menurut kriteria pada Tabel 3.5.",
        "Analisis tingkat penyelesaian tugas: persentase responden yang "
        "berhasil menyelesaikan setiap tugas tanpa bantuan.",
        "Analisis tematik: hasil wawancara diringkas menurut tema untuk "
        "merumuskan saran perbaikan dan pengembangan lanjut.",
    ], numbering="number")

    add_table(doc,
              ["Rentang Skor SUS", "Kategori", "Tafsiran"],
              [["\u2265 80,3", "Grade A", "Excellent \u2014 sangat baik"],
               ["68 \u2013 80,2", "Grade B\u2013C", "Good / Acceptable \u2014 "
                                  "dapat diterima"],
               ["51 \u2013 67,9", "Grade D", "Marginal \u2014 perlu "
                                  "perbaikan"],
               ["< 51", "Grade F", "Not acceptable \u2014 tidak dapat "
                                   "diterima"]],
              widths=[3.5, 3.5, 7.0], fs=10)
    caption(doc, "Tabel 3.5 Kriteria penilaian skor System Usability Scale "
                 "(SUS)")

    h2(doc, "H. Jadwal Penelitian")
    body(doc, "Pelaksanaan penelitian direncanakan berlangsung selama enam "
              "bulan dengan rincian sebagaimana disajikan pada Tabel 3.6.",
         indent=True, space_after=8)
    add_table(doc,
              ["Bulan ke-", "Kegiatan", "Luaran"],
              [["1", "Studi pendahuluan, observasi, wawancara, dan studi "
                     "pustaka", "Laporan analisis kebutuhan"],
               ["2", "Perancangan use case, ERD, arsitektur, kontrak API, "
                     "dan antarmuka", "Dokumen perancangan dan proposal"],
               ["3", "Implementasi modul autentikasi, katalog, dan ruang "
                     "baca digital", "Purwarupa pertama"],
               ["4", "Implementasi transaksi QRIS, penyelesaian idempoten, "
                     "dan dasbor pengelola", "Purwarupa lengkap"],
               ["5", "Pengujian internal (black box dan eksperimen "
                     "idempotensi) serta perbaikan", "Laporan hasil "
                                                     "pengujian"],
               ["6", "Uji coba pengguna, pengukuran SUS, evaluasi, dan "
                     "penyusunan laporan akhir", "Laporan skripsi dan produk "
                                                 "akhir"]],
              widths=[2.0, 7.5, 4.5], fs=10)
    caption(doc, "Tabel 3.6 Jadwal pelaksanaan penelitian")


# ======================== DAFTAR PUSTAKA ========================

def build_daftar_pustaka(doc):
    h1(doc, "DAFTAR PUSTAKA")
    body(doc, "Daftar pustaka berikut memuat sumber primer dan sekunder "
              "yang relevan dengan topik penelitian. Seluruh sumber jurnal "
              "merupakan terbitan tahun 2021\u20132025, terindeks dalam basis "
              "data SINTA (Science and Technology Index) atau terindeks "
              "secara internasional, serta memiliki Digital Object Identifier "
              "(DOI) yang dapat diakses melalui pranala yang dicantumkan.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=11, italic=True,
         line_spacing=1.5, space_after=12)

    with open(DATA_PUSTAKA, "r", encoding="utf-8") as f:
        refs = json.load(f)

    refs_sorted = sorted(
        refs,
        key=lambda r: r["authors"].replace("& ", "").split(",")[0].strip().lower()
    )

    for r in refs_sorted:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = 1.5
        pf.first_line_indent = Cm(-1.27)
        pf.left_indent = Cm(1.27)
        pf.space_after = Pt(8)

        if r["type"] == "buku":
            txt = (f"{r['authors']} ({r['year']}). {r['title']}. "
                   f"{r['city']}: {r['publisher']}.")
            run = p.add_run(txt)
            run.font.size = Pt(12)
            set_run_font(run)
        else:
            vol = r.get("volume", "")
            pg = r.get("pages", "")
            tail = ", ".join([x for x in (vol, pg) if x])
            base = f"{r['authors']} ({r['year']}). {r['title']}. {r['journal']}"
            if tail:
                base += f", {tail}"
            base += ". "
            run = p.add_run(base)
            run.font.size = Pt(12)
            set_run_font(run)
            run2 = p.add_run("https://doi.org/")
            run2.font.size = Pt(12)
            set_run_font(run2)
            add_hyperlink(p, r["url"], r["doi"])


# ======================== LAMPIRAN ========================

def build_lampiran(doc):
    h1(doc, "LAMPIRAN")

    h2(doc, "Lampiran 1. Kuesioner System Usability Scale (SUS)")
    body(doc, "Petunjuk: berilah tanda centang (\u2713) pada kolom yang "
              "paling sesuai dengan pendapat Anda setelah menggunakan "
              "platform Wahidiyah Book. Pilihan: 1 = Sangat Tidak Setuju "
              "(STS), 2 = Tidak Setuju (TS), 3 = Netral (N), 4 = Setuju (S), "
              "5 = Sangat Setuju (SS).")
    add_table(doc,
              ["No.", "Pernyataan", "1", "2", "3", "4", "5"],
              [[f"{i}.", t, "", "", "", "", ""] for i, t in enumerate([
                  "Saya akan menggunakan platform ini lagi di lain waktu.",
                  "Saya merasa platform ini rumit untuk digunakan.",
                  "Saya merasa platform ini mudah digunakan.",
                  "Saya memerlukan bantuan orang lain untuk menggunakan "
                  "platform ini.",
                  "Fitur-fitur pada platform ini terasa terpadu dengan baik.",
                  "Menurut saya, banyak hal yang tidak konsisten pada "
                  "platform ini.",
                  "Saya rasa kebanyakan orang dapat mempelajari platform "
                  "dengan cepat.",
                  "Saya merasa platform ini membingungkan.",
                  "Saya tidak menemukan kendala berarti saat menggunakan "
                  "platform ini.",
                  "Saya perlu membiasakan diri terlebih dahulu sebelum "
                  "dapat menggunakan platform ini.",
              ], start=1)],
              widths=[1.0, 8.5, 1.0, 1.0, 1.0, 1.0, 1.0], fs=9.5)
    add_paragraph(doc, "", size=6, line_spacing=1.0, space_after=8)
    body(doc, "Cara penskoran: untuk butir bernomor ganjil, kurangi nilainya "
              "dengan 1; untuk butir bernomor genap, kurangi 5 dengan "
              "nilainya. Jumlahkan seluruh hasilnya, lalu kalikan dengan 2,5 "
              "untuk memperoleh skor SUS dalam rentang 0\u2013100.",
         italic=True)

    h2(doc, "Lampiran 2. Pedoman Wawancara Pengelola dan Penilai Ahli")
    bullets(doc, [
        "Bagaimana alur pengelolaan bahan bacaan yang berjalan saat ini?",
        "Kendala apa yang paling sering muncul dalam penelusuran dan "
        "peminjaman buku?",
        "Bagaimana pendapat Bapak/Ibu mengenai penerapan langganan berbayar "
        "untuk ruang baca digital?",
        "Apakah alur pembayaran QRIS yang dirancang sudah sesuai dengan "
        "kebiasaan pengguna?",
        "Apakah pembatasan akses konten berbayar sudah berjalan sebagaimana "
        "mestinya?",
        "Bagaimana penilaian Bapak/Ibu terhadap keamanan notifikasi "
        "pembayaran yang diterapkan?",
        "Fitur atau perbaikan apa yang perlu ditambahkan pada rilis "
        "berikutnya?",
    ], numbering="number")

    h2(doc, "Lampiran 3. Cuplikan Kode Inti Mekanisme Penyelesaian "
            "Transaksi")
    body(doc, "Cuplikan berikut menunjukkan inti mekanisme penyelesaian "
              "transaksi yang idempoten: pembaruan bersyarat sebagai "
              "compare-and-swap, pemeriksaan jumlah baris yang terpengaruh, "
              "serta transaksi kompensasi apabila aktivasi langganan gagal.",
         indent=True, space_after=6)
    code_block(doc, [
        "// server/routes/payments.routes.js",
        "async function settleTransaction(transaction) {",
        "  if (!transaction || transaction.status === 'success') {",
        "    return { alreadyProcessed: true, activated: null };",
        "  }",
        "",
        "  // --- compare-and-swap: klaim atomik atas keadaan 'pending' ---",
        "  const { data: claimed, error: claimError } = await supabaseServer",
        "    .from('transactions')",
        "    .update({ status: 'success', verified_at: new Date().toISOString() })",
        "    .eq('id', transaction.id)",
        "    .eq('status', TRANSACTION_PENDING_STATUS)   // syarat keadaan",
        "    .select('id');",
        "",
        "  if (claimError) throw claimError;",
        "  if (!claimed || claimed.length === 0) {",
        "    // kalah klaim: pekerja lain sudah menyelesaikan -> idempoten",
        "    return { alreadyProcessed: true, activated: null };",
        "  }",
        "",
        "  let activated;",
        "  try {",
        "    activated = await activatePaidMembership({",
        "      userId: transaction.user_id,",
        "      planName: transaction.plan_name,",
        "      price: transaction.amount,",
        "      refCode: toRefCode(transaction.ref_no)",
        "    });",
        "  } catch (activationError) {",
        "    // --- transaksi kompensasi ---",
        "    await supabaseServer",
        "      .from('transactions')",
        "      .update({ status: TRANSACTION_PENDING_STATUS, verified_at: null })",
        "      .eq('id', transaction.id);",
        "    throw activationError;",
        "  }",
        "",
        "  return { alreadyProcessed: false, activated };",
        "}",
    ])
    add_paragraph(doc, "", size=6, line_spacing=1.0, space_after=8)

    body(doc, "Cuplikan berikut menunjukkan verifikasi notifikasi yang "
              "bersifat fail-closed dengan tanda tangan SHA-512 dan "
              "perbandingan waktu konstan.", indent=True, space_after=6)
    code_block(doc, [
        "// server/lib/midtrans.js",
        "export function verifySignatureKey(payload = {}) {",
        "  if (!MIDTRANS_SERVER_KEY) {",
        "    // gagal -> TUTUP, bukan lewat (fail-closed)",
        "    return { ok: false, reason: 'MIDTRANS_SERVER_KEY belum diisi.' };",
        "  }",
        "",
        "  const raw = `${payload.order_id ?? ''}${payload.status_code ?? ''}",
        "               ${payload.gross_amount ?? ''}${MIDTRANS_SERVER_KEY}`;",
        "  const expected = crypto.createHash('sha512').update(raw).digest('hex');",
        "  const provided = Buffer.from(String(payload.signature_key ?? ''));",
        "  const expectedBuffer = Buffer.from(expected);",
        "",
        "  if (provided.length !== expectedBuffer.length) {",
        "    return { ok: false, reason: 'Signature key tidak valid.' };",
        "  }",
        "",
        "  return crypto.timingSafeEqual(provided, expectedBuffer)",
        "    ? { ok: true }",
        "    : { ok: false, reason: 'Signature key tidak valid.' };",
        "}",
    ])
    add_paragraph(doc, "", size=6, line_spacing=1.0, space_after=8)

    body(doc, "Cuplikan berikut menunjukkan kebijakan tembolok selektif "
              "pada service worker, yang secara sadar melewati konten "
              "berbayar, autentikasi, langganan, dan transaksi.",
         indent=True, space_after=6)
    code_block(doc, [
        "// public/sw.js",
        "self.addEventListener('fetch', (event) => {",
        "  const { request } = event;",
        "  const url = new URL(request.url);",
        "",
        "  if (request.method !== 'GET') return;",
        "",
        "  // JANGAN cache: konten buku berbayar, autentikasi, langganan,",
        "  // dan transaksi (Invarian Keamanan Konten)",
        "  if (url.pathname.includes('/api/books/') &&",
        "      url.pathname.endsWith('/content') ||",
        "      url.pathname.startsWith('/api/auth/') ||",
        "      url.pathname.startsWith('/api/subscriptions/') ||",
        "      url.pathname.startsWith('/api/transactions')) {",
        "    return;",
        "  }",
        "",
        "  // Data publik (katalog, agenda, iklan): Network First +",
        "  // fallback ke cache",
        "  // ...",
        "});",
    ])


# ======================== MAIN PROPOSAL ========================

def build_proposal():
    doc = Document()
    set_default_font(doc)
    _init_heading_styles(doc)

    s1 = doc.sections[0]
    set_page_margins(s1)

    build_cover(doc)
    add_page_break(doc)
    build_halaman_judul(doc)

    s2 = add_section_break(doc)
    set_page_number_format(s2, fmt='lowerRoman', start=1)
    set_footer_page_number(s2)

    add_page_break(doc)
    build_persetujuan(doc)
    add_page_break(doc)
    build_pengesahan(doc)
    add_page_break(doc)
    build_surat_pernyataan(doc)
    add_page_break(doc)
    build_kata_pengantar(doc)
    add_page_break(doc)
    build_abstrak(doc)
    add_page_break(doc)
    build_daftar_isi(doc)
    add_page_break(doc)
    build_daftar_tabel(doc)
    add_page_break(doc)
    build_daftar_gambar(doc)

    s3 = add_section_break(doc)
    set_page_number_format(s3, fmt='decimal', start=1)
    set_footer_page_number(s3)

    build_bab1(doc)
    add_page_break(doc)
    build_bab2(doc)
    add_page_break(doc)
    build_bab3(doc)
    add_page_break(doc)
    build_daftar_pustaka(doc)
    add_page_break(doc)
    build_lampiran(doc)

    # minta Word memperbarui field (daftar isi) saat dokumen dibuka
    settings = doc.settings.element
    uf = OxmlElement('w:updateFields')
    uf.set(qn('w:val'), 'true')
    settings.append(uf)

    out = os.path.join(HERE, "Proposal_Skripsi_Rancang_Bangun_20232210156.docx")
    doc.save(out)
    print(f"[OK] Proposal -> {out} ({os.path.getsize(out):,} bytes)")
    return out


# ======================== FORM PENGAJUAN JUDUL ========================

def build_pengajuan_judul():
    doc = Document()
    set_default_font(doc)
    _init_heading_styles(doc)
    sec = doc.sections[0]
    set_page_margins(sec)
    sec.top_margin = Cm(3)
    sec.left_margin = Cm(3)
    sec.right_margin = Cm(3)
    sec.bottom_margin = Cm(3)

    add_paragraph(doc, UNIV, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                  size=12, line_spacing=1.0, space_before=4, space_after=2)
    add_paragraph(doc, FAKULTAS, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                  size=11, line_spacing=1.0, space_after=2)
    add_paragraph(doc, PRODI, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                  size=11, line_spacing=1.0, space_after=10)
    if os.path.exists(LOGO_PATH):
        add_centered_image(doc, LOGO_PATH, 3.5)

    add_paragraph(doc, "PENGAJUAN JUDUL SKRIPSI", bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, size=14,
                  line_spacing=1.5, space_before=16, space_after=16)

    add_paragraph(doc, "Yang bertanda tangan di bawah ini:",
                  align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=2.0,
                  first_line_indent=Cm(0.5), space_after=6)
    add_paragraph(doc, f"Nama\t\t: {NAMA}", line_spacing=2.0, space_after=4,
                  first_line_indent=Cm(0.5))
    add_paragraph(doc, f"NIM\t\t: {NIM}", line_spacing=2.0, space_after=4,
                  first_line_indent=Cm(0.5))
    add_paragraph(doc, "Program Studi\t: Teknik Informatika",
                  line_spacing=2.0, space_after=12, first_line_indent=Cm(0.5))

    add_paragraph(doc, "Mengajukan judul skripsi sebagai berikut:",
                  align=WD_ALIGN_PARAGRAPH.LEFT, line_spacing=2.0,
                  first_line_indent=Cm(0.5), space_after=6)

    for _ in range(5):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(2)
        pf.line_spacing = 2.0
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("_" * 92)
        r.font.size = Pt(11)
        set_run_font(r)
    add_paragraph(doc, "", size=11, line_spacing=1.0, space_after=6)

    add_paragraph(doc, "Dengan obyek penelitian pada platform Wahidiyah Book "
                       "di lingkungan Universitas Wahidiyah Kediri.",
                  line_spacing=2.0, first_line_indent=Cm(0.5),
                  space_after=10, italic=True)

    add_paragraph(doc, "Sebagai pertimbangan, berikut disampaikan rancangan "
                       "penelitian:", line_spacing=2.0,
                  first_line_indent=Cm(0.5), space_after=4)
    for i, t in enumerate([
        "Jenis penelitian: Research and Development (rancang bangun) dengan "
        "model prototyping.",
        "Obyek yang dikembangkan: platform web katalog buku dan ruang baca "
        "digital berlangganan yang terintegrasi pembayaran QRIS (React 18 + "
        "Vite 6, Express 4, PostgreSQL/Supabase, Midtrans Core API).",
        "Fokus kebaruan: mekanisme penyelesaian transaksi yang idempoten "
        "berbasis compare-and-swap atomik dengan transaksi kompensasi, "
        "verifikasi notifikasi fail-closed (SHA-512 + timingSafeEqual), dan "
        "dua jalur penyelesaian (webhook dan polling otoritatif server).",
        "Desain uji coba: pengujian black box berbasis ISO/IEC 25010, "
        "eksperimen idempotensi dengan pengiriman notifikasi ganda, dan "
        "kuesioner System Usability Scale (SUS).",
        "Subjek uji coba: pengelola (1\u20132 orang), mahasiswa pengguna "
        "(\u2265 30 orang), dan penilai ahli (2\u20133 orang).",
        "Luaran yang diharapkan: platform berfungsi penuh tanpa aktivasi "
        "langganan ganda, dengan skor SUS \u2265 68.",
    ], start=1):
        add_paragraph(doc, f"{i}. {t}", line_spacing=2.0,
                      first_line_indent=Cm(0.5), space_after=4,
                      align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    add_paragraph(doc,
                  "Demikian pengajuan judul skripsi ini disampaikan. Atas "
                  "perhatian dan kesediaan Bapak/Ibu untuk memeriksa, "
                  "disampaikan terima kasih teriring do\u2019a "
                  "\u201CJaz\u0101kumull\u0101hu khairan wa sa\u2019\u0101datid "
                  "dun-y\u0101 wal \u0101khirah\u201D. \u0100m\u012bn.",
                  line_spacing=2.0, first_line_indent=Cm(0.5),
                  space_before=12, space_after=12,
                  align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    add_paragraph(doc, f"{KOTA}, .......................... {TAHUN}",
                  align=WD_ALIGN_PARAGRAPH.RIGHT, line_spacing=1.5,
                  space_after=6)
    add_paragraph(doc, "Mengetahui,\t\t\t\t\tMenyetujui,", bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.5,
                  space_after=2)
    add_paragraph(doc, "Ketua Program Studi\t\t\t\tPembimbing,",
                  align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.5,
                  space_after=48)
    add_paragraph(doc, "(........................................)\t\t\t\t"
                       "(........................................)",
                  align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0,
                  space_after=2)
    add_paragraph(doc, "NIY. ..............................\t\t\t\t"
                       "NIY. ..............................",
                  align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.0,
                  space_after=16)

    add_paragraph(doc, "Keterangan Judul Penelitian: ACC / TIDAK ACC dengan "
                       "perubahan:", align=WD_ALIGN_PARAGRAPH.LEFT,
                  line_spacing=1.5, space_after=4, first_line_indent=Cm(0.5))
    for _ in range(4):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run("_" * 100)
        r.font.size = Pt(11)
        set_run_font(r)

    out = os.path.join(HERE, "Pengajuan_Judul_Skripsi_20232210156.docx")
    doc.save(out)
    print(f"[OK] Pengajuan judul -> {out} ({os.path.getsize(out):,} bytes)")
    return out


# ======================== FORM KERANGKA MODEL PENELITIAN ========================

def build_kerangka_model():
    doc = Document()
    set_default_font(doc)
    _init_heading_styles(doc)
    sec = doc.sections[0]
    set_page_margins(sec)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Cm(29.7)
    sec.page_height = Cm(21.0)
    sec.top_margin = Cm(2)
    sec.left_margin = Cm(2)
    sec.right_margin = Cm(2)
    sec.bottom_margin = Cm(2)

    if os.path.exists(LOGO_PATH):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(6)
        p.add_run().add_picture(LOGO_PATH, width=Cm(2.5))

    add_paragraph(doc, "KERANGKA MODEL PENELITIAN", bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, size=14,
                  line_spacing=1.5, space_before=0, space_after=2,
                  underline=True)
    add_paragraph(doc, f"Nama\t\t: {NAMA}", line_spacing=1.5,
                  first_line_indent=Cm(0.5), space_after=2)
    add_paragraph(doc, f"NIM\t\t: {NIM}", line_spacing=1.5,
                  first_line_indent=Cm(0.5), space_after=2)
    add_paragraph(doc, "Program Studi\t: Teknik Informatika",
                  line_spacing=1.5, first_line_indent=Cm(0.5), space_after=2)
    add_paragraph(doc, f"Fakultas\t\t: {FAKULTAS.title()}",
                  line_spacing=1.5, first_line_indent=Cm(0.5), space_after=10)

    add_paragraph(doc, "Topik/Judul:", bold=True, line_spacing=1.5,
                  first_line_indent=Cm(0.5), space_after=4)
    add_paragraph(doc, f"\u201C{JUDUL}\u201D", bold=True,
                  align=WD_ALIGN_PARAGRAPH.CENTER, line_spacing=1.5,
                  space_before=4, space_after=12)

    headers = ["Latar Belakang", "Masalah", "Definisi Operasional Variabel",
               "Metode", "Langkah-langkah", "Instrumen Penelitian"]

    rows = [(
        "1. Pengelolaan katalog dan ruang baca di lingkungan Universitas "
        "Wahidiyah Kediri masih terpisah-pisah dan belum terintegrasi dengan "
        "pembayaran digital.\n\n"
        "2. Pengguna kesulitan menemukan bahan bacaan; belum tersedia ruang "
        "baca digital berlangganan.\n\n"
        "3. Notifikasi payment gateway bersifat at-least-once dan dapat "
        "hilang, sehingga berisiko menimbulkan aktivasi langganan ganda "
        "atau pembayaran yang tidak tercatat.\n\n"
        "4. Pengembangan lokal tanpa URL publik menyulitkan pengujian alur "
        "pembayaran.\n\n"
        "5. Belum ada platform sejenis di lingkungan Universitas Wahidiyah "
        "Kediri, sehingga kebutuhan perlu dikristalkan melalui purwarupa.",

        "1. Bagaimana merancang dan membangun platform Wahidiyah Book yang "
        "mengintegrasikan katalog, ruang baca digital, dan langganan "
        "berbayar?\n"
        "2. Bagaimana merancang penyelesaian transaksi QRIS yang "
        "idempoten?\n"
        "3. Bagaimana merancang verifikasi notifikasi yang fail-closed?\n"
        "4. Bagaimana menyediakan dua jalur penyelesaian transaksi?\n"
        "5. Bagaimana mencegah terbentuknya akun yang belum terverifikasi?\n"
        "6. Bagaimana mencegah tersimpannya konten berbayar di tembolok "
        "PWA?\n"
        "7. Bagaimana kualitas fungsional dan kebergunaan produk?",

        "1. Penyelesaian transaksi idempoten: keadaan akhir sama berapa pun "
        "kali operasi dijalankan untuk satu order_id.\n"
        "2. Rasio aktivasi ganda: (aktivasi ganda / seluruh order_id) \u00D7 "
        "100%; target 0%.\n"
        "3. Verifikasi fail-closed: notifikasi tanpa tanda tangan sah ditolak "
        "dengan HTTP 401.\n"
        "4. Dua jalur penyelesaian: webhook dan pemeriksaan status "
        "otoritatif server.\n"
        "5. Registrasi dua fase: baris users dibuat hanya setelah OTP sah.\n"
        "6. Tembolok selektif PWA: konten berbayar, autentikasi, langganan, "
        "dan transaksi dilewati dari tembolok.\n"
        "7. Kesesuaian fungsional (ISO/IEC 25010) dan kebergunaan (SUS).",

        "1. Jenis: Research and Development (rancang bangun) dengan model "
        "prototyping.\n"
        "2. Tempat: lingkungan Universitas Wahidiyah Kediri.\n"
        "3. Pendekatan: arsitektur tiga lapis, REST + JSON, SPA dengan "
        "React 18 + Vite 6, basis data PostgreSQL (Supabase).\n"
        "4. Integrasi: Midtrans Core API (QRIS) dan Gmail API.\n"
        "5. Subjek uji coba: pengelola (1\u20132), mahasiswa pengguna "
        "(\u2265 30), penilai ahli (2\u20133).\n"
        "6. Desain uji coba: one-shot post-test only.\n"
        "7. Teknik pengumpulan data: skenario black box, log server, "
        "kuesioner SUS, wawancara.",

        "1. Studi pendahuluan: observasi, wawancara, studi pustaka.\n"
        "2. Perancangan: use case, ERD, arsitektur, kontrak API, "
        "antarmuka.\n"
        "3. Implementasi: klien React, server Express, skema basis data, "
        "mekanisme penyelesaian idempoten.\n"
        "4. Pengujian internal: black box berbasis ISO/IEC 25010 dan "
        "eksperimen idempotensi.\n"
        "5. Uji coba pengguna: skenario tugas dan kuesioner SUS.\n"
        "6. Evaluasi dan penyempurnaan iteratif sampai spesifikasi "
        "terpenuhi.",

        "1. Lembar skenario pengujian black box.\n"
        "2. Skrip permintaan berulang untuk uji idempotensi.\n"
        "3. Log server dan catatan basis data.\n"
        "4. Kuesioner System Usability Scale (10 butir).\n"
        "5. Lembar observasi penyelesaian tugas.\n"
        "6. Pedoman wawancara pengelola dan penilai ahli.\n"
        "7. Dokumentasi kode sumber dan skema basis data."
    )]

    add_table(doc, headers, rows, fs=9, hdr_fs=9.5)

    add_paragraph(doc, "", line_spacing=1.0, space_after=6)
    add_paragraph(doc, f"{KOTA}, .......................... {TAHUN}",
                  align=WD_ALIGN_PARAGRAPH.RIGHT, line_spacing=1.5,
                  space_before=18, space_after=6)
    add_paragraph(doc, "Mahasiswa Peneliti,", align=WD_ALIGN_PARAGRAPH.RIGHT,
                  line_spacing=1.5, space_after=48)
    add_paragraph(doc, NAMA, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT,
                  line_spacing=1.0, space_after=2)
    add_paragraph(doc, f"NIM. {NIM}", align=WD_ALIGN_PARAGRAPH.RIGHT,
                  line_spacing=1.0, space_after=6)

    out = os.path.join(HERE, "Kerangka_Model_Penelitian_20232210156.docx")
    doc.save(out)
    print(f"[OK] Kerangka model -> {out} ({os.path.getsize(out):,} bytes)")
    return out


if __name__ == "__main__":
    build_proposal()
    build_pengajuan_judul()
    build_kerangka_model()
    print("\n=== BERKAS TERGENERASI ===")
    for f in sorted(os.listdir(HERE)):
        if f.endswith(".docx"):
            print(f"  - {f}  ({os.path.getsize(os.path.join(HERE, f)):,} bytes)")
